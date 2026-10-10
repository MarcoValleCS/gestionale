"""Viste di vendite: preventivi e ordini cliente."""
from decimal import Decimal

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db.models import Count, Q
from django.forms import modelformset_factory
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, DetailView, ListView, UpdateView

from apps.accounts.permissions import ROLE_ADMIN, ROLE_PURCHASING, ROLE_SALES, ROLE_WAREHOUSE, RoleRequiredMixin, role_required
from apps.core.concurrency import ConflictAwareUpdateView
from apps.contacts.models import Contact
from apps.catalog.models import Product
from apps.core.models import VatRate

from . import analytics, emailing, services
from .forms import (
    QuoteForm,
    QuoteLineForm,
    QuoteTemplateForm,
    QuoteTemplateLineForm,
    SalesOrderForm,
    SalesOrderLineForm,
)
from .models import (
    LINE_ARTICLE,
    Quote,
    QuoteLine,
    QuoteTemplate,
    QuoteTemplateLine,
    SalesOrder,
    SalesOrderLine,
    group_lines_by_section,
)

QuoteLineFormSet = modelformset_factory(QuoteLine, form=QuoteLineForm, extra=0, can_delete=True)
SalesOrderLineFormSet = modelformset_factory(SalesOrderLine, form=SalesOrderLineForm, extra=0, can_delete=True)
QuoteTemplateLineFormSet = modelformset_factory(QuoteTemplateLine, form=QuoteTemplateLineForm, extra=0, can_delete=True)

QUOTE_ROLES = (ROLE_ADMIN, ROLE_SALES)
ORDER_VIEW_ROLES = (ROLE_ADMIN, ROLE_SALES, ROLE_WAREHOUSE, ROLE_PURCHASING)
ORDER_EDIT_ROLES = (ROLE_ADMIN, ROLE_SALES)
ORDER_DELIVER_ROLES = (ROLE_ADMIN, ROLE_SALES, ROLE_WAREHOUSE)

MESI_ITALIANI = [
    "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
    "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre",
]


# ------------------------------------------------------------------ helper
def fdate(value):
    return value.strftime("%d/%m/%Y") if value else ""


def statistics(request):
    """Statistiche complete: fatturato, margine e classifiche del periodo."""
    from apps.core.cache import memoizza

    period = request.GET.get("periodo", analytics.PERIOD_YEAR)
    if period not in {value for value, _label in analytics.PERIOD_CHOICES}:
        period = analytics.PERIOD_YEAR

    try:
        months = int(request.GET.get("finestra", 12))
    except (TypeError, ValueError):
        months = 12
    if months not in (1, 3, 6, 12):
        months = 12
    try:
        offset = int(request.GET.get("indietro", 0))
    except (TypeError, ValueError):
        offset = 0
    offset = max(0, min(offset, 36))

    stats = memoizza(f"analytics_breakdowns_{period}", lambda: analytics.breakdowns(period, limit=None), 60)
    return render(
        request,
        "sales/statistics.html",
        {
            "page_title": "Statistiche",
            "period": period,
            "period_choices": analytics.PERIOD_CHOICES,
            "stats": stats["summary"],
            "margin_by_product": stats["by_product"],
            "margin_by_supplier": stats["by_supplier"],
            "margin_by_customer": stats["by_customer"],
            "margin_by_job": stats["by_job"],
            "series": memoizza(
                f"analytics_serie_{months}_{offset}", lambda: analytics.monthly_series(months, end_offset=offset), 60
            ),
            "months": months,
            "offset": offset,
            "months_choices": (1, 3, 6, 12),
        },
    )


@role_required(*QUOTE_ROLES)
def follow_up(request):
    """Preventivi da sollecitare e tasso di conversione."""
    from datetime import timedelta

    oggi = timezone.localdate()
    try:
        giorni = int(request.GET.get("giorni", 7))
    except (TypeError, ValueError):
        giorni = 7
    giorni = max(1, min(giorni, 365))

    aperti = list(
        Quote.objects.filter(status__in=[Quote.STATUS_DRAFT, Quote.STATUS_SENT])
        .select_related("customer", "created_by")
        .order_by("date")
    )
    da_sollecitare = []
    for preventivo in aperti:
        attesa = (oggi - preventivo.date).days
        scadenza = preventivo.valid_until
        giorni_a_scadenza = (scadenza - oggi).days if scadenza else None
        preventivo.attesa = attesa
        preventivo.giorni_a_scadenza = giorni_a_scadenza
        preventivo.scaduto = giorni_a_scadenza is not None and giorni_a_scadenza < 0
        preventivo.email_cliente = bool(preventivo.customer.email)
        if attesa >= giorni:
            da_sollecitare.append(preventivo)

    # conversione: quanti preventivi diventano ordini
    inviati = Quote.objects.filter(status__in=[Quote.STATUS_SENT, Quote.STATUS_ACCEPTED, Quote.STATUS_CONVERTED]).count()
    convertiti = Quote.objects.filter(status=Quote.STATUS_CONVERTED).count()
    rifiutati = Quote.objects.filter(status=Quote.STATUS_REJECTED).count()
    conclusi = convertiti + rifiutati
    conversione = round(convertiti / conclusi * 100, 1) if conclusi else 0

    # tempo medio di chiusura dei preventivi convertiti
    chiusure = []
    for preventivo in Quote.objects.filter(status=Quote.STATUS_CONVERTED).prefetch_related("generated_orders")[:200]:
        for ordine in preventivo.generated_orders.all():
            delta = (ordine.date - preventivo.date).days
            if delta >= 0:
                chiusure.append(delta)
    tempo_medio = round(sum(chiusure) / len(chiusure), 1) if chiusure else None

    valore_fermo = sum((p.grand_total or 0) for p in da_sollecitare)

    return render(
        request,
        "sales/follow_up.html",
        {
            "page_title": "Preventivi da seguire",
            "da_sollecitare": da_sollecitare,
            "giorni": giorni,
            "inviati": inviati,
            "convertiti": convertiti,
            "rifiutati": rifiutati,
            "conversione": conversione,
            "tempo_medio": tempo_medio,
            "valore_fermo": valore_fermo,
            "email_configurata": emailing.email_configured(),
        },
    )


@role_required(*QUOTE_ROLES)
def quote_reminder(request, pk):
    """Invia un sollecito per un preventivo senza risposta."""
    preventivo = get_object_or_404(Quote.objects.select_related("customer"), pk=pk)
    if request.method != "POST":
        return redirect("sales:follow_up")
    if not preventivo.customer.email:
        messages.error(request, f"Il cliente {preventivo.customer.name} non ha un indirizzo email.")
        return redirect("sales:follow_up")
    if not emailing.email_configured():
        messages.error(request, "Invio email non configurato: imposta le variabili EMAIL_* nel file .env.")
        return redirect("sales:follow_up")

    giorni = (timezone.localdate() - preventivo.date).days
    from apps.core.email_templates import contenuto
    from apps.core.mailing import accoda_email
    from apps.core.models import CompanySettings, EmailInCoda, EmailTemplate
    from apps.core.utils import format_money

    oggetto, corpo = contenuto(
        EmailTemplate.KIND_QUOTE_REMINDER,
        {
            "cliente": preventivo.customer.name,
            "numero": preventivo.number,
            "data": f"{preventivo.date:%d/%m/%Y}",
            "totale": format_money(preventivo.grand_total),
            "giorni": giorni,
            "azienda": CompanySettings.load().name,
            "termini": (
                f"Condizioni generali di vendita: {request.build_absolute_uri(reverse('core:termini'))}\n"
            ),
        },
    )
    try:
        pdf = emailing.render_quote_pdf(preventivo)
        accoda_email(
            to_email=preventivo.customer.email,
            subject=oggetto,
            message=corpo,
            attachment=pdf,
            attachment_name=f"Preventivo_{preventivo.number}.pdf",
            descrizione=f"Sollecito {preventivo.number} a {preventivo.customer.email}",
            modello=EmailInCoda.MODELLO_PREVENTIVO,
            oggetto_id=preventivo.pk,
        )
    except Exception as exc:
        messages.error(request, f"Sollecito non accodato: {exc}")
    else:
        messages.success(request, f"Sollecito accodato per {preventivo.customer.email} ({preventivo.number}): verrà inviato entro pochi minuti.")
    return redirect("sales:follow_up")


@role_required(*QUOTE_ROLES)
def commission_report(request):
    """Resoconto mensile delle provvigioni da riconoscere e a chi."""
    oggi = timezone.localdate()
    try:
        year = int(request.GET.get("anno", oggi.year))
    except (TypeError, ValueError):
        year = oggi.year
    try:
        month = int(request.GET.get("mese", oggi.month))
    except (TypeError, ValueError):
        month = oggi.month
    if not 1 <= month <= 12:
        month = oggi.month
    if not 2000 <= year <= 2100:
        year = oggi.year

    stati = request.GET.get("stato", "tutti")
    stati_disponibili = {
        "tutti": ("Tutti gli ordini", analytics.STATO_TUTTI),
        "consegnati": ("Solo consegnati", analytics.STATO_CONSEGNATI),
        "confermati": ("Solo in lavorazione", analytics.STATO_CONFERMATI),
    }
    if stati not in stati_disponibili:
        stati = "tutti"

    precedente = analytics.add_months(year, month, -1)
    successivo = analytics.add_months(year, month, 1)

    return render(
        request,
        "sales/commission_report.html",
        {
            "page_title": "Provvigioni",
            "report": analytics.commission_report(year, month, stati_disponibili[stati][1]),
            "riepilogo_anno": analytics.commission_year(year, stati_disponibili[stati][1]),
            "stato": stati,
            "stati_disponibili": [(chiave, etichetta) for chiave, (etichetta, _s) in stati_disponibili.items()],
            "anno_corrente": year,
            "anni": list(range(oggi.year - 3, oggi.year + 2)),
            "mesi": [(numero, nome) for numero, nome in enumerate(MESI_ITALIANI, start=1)],
            "precedente": {"anno": precedente[0], "mese": precedente[1]},
            "successivo": {"anno": successivo[0], "mese": successivo[1]},
        },
    )


def plain_number(value):
    """Numero decimale senza zeri finali, per i campi dei form via JavaScript."""
    text = f"{Decimal(value or 0):f}"
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def build_print_context(
    document,
    *,
    title,
    counterparty,
    meta_rows,
    back_url,
    counterparty_label="Destinatario",
    notes="",
    show_signature=False,
    signature_label="",
    show_prices=True,
    show_discount=True,
    extra_fields=None,
    gross_amounts=False,
):
    """Contesto per il documento stampabile (templates/print/document.html).

    Con ``gross_amounts=True`` gli importi di riga e i totali di sezione
    includono l'IVA (stampe al cliente: preventivi e conferme d'ordine).
    """
    lines = document.lines.select_related("product", "uom", "vat_rate")
    # numero di colonne della tabella righe: serve ai colspan delle sezioni.
    # Viene calcolato qui per non lasciare aritmetica fragile nel template.
    if show_prices:
        colonne = 8 if show_discount else 7
    else:
        colonne = 4
    return {
        "document": document,
        "document_title": title,
        "counterparty": counterparty,
        "counterparty_label": counterparty_label,
        "meta_rows": meta_rows,
        "extra_fields": extra_fields or [],
        "back_url": back_url,
        "lines": lines,
        "line_groups": group_lines_by_section(lines),
        "vat_rows": document.vat_breakdown() if show_prices else [],
        "notes": notes,
        "show_signature": show_signature,
        "signature_label": signature_label,
        "show_prices": show_prices,
        "show_discount": show_discount,
        "gross_amounts": gross_amounts,
        "table_columns": colonne,
        # Le viste di stampa lo valorizzano (download del PDF generato dal
        # gestionale); resta vuoto per gli altri documenti (fatture, DDT…).
        "download_url": "",
    }


def with_vat_rates(context, documento=None):
    from apps.core.forms import active_units, active_vat_rates

    context["vat_rates_json"] = {str(v.pk): str(v.rate) for v in VatRate.objects.filter(is_active=True)}
    context["quick_uoms"] = active_units()
    context["quick_vats"] = active_vat_rates()
    context["quick_suppliers"] = Contact.objects.filter(is_supplier=True, active=True).order_by("name")
    context["customer_discounts_json"] = {
        str(contact.pk): str(contact.sale_discount_pct)
        for contact in Contact.objects.filter(is_customer=True, active=True, sale_discount_pct__gt=0)
    }
    # Per il margine servono i costi: si caricano solo quelli degli articoli già
    # presenti nel documento (pochi). Caricare tutto il catalogo significherebbe
    # un JSON enorme da scaricare a ogni apertura del modulo. Gli articoli scelti
    # durante la compilazione portano il costo con la chiamata di autocompletamento.
    ids = []
    if documento is not None:
        ids = [pk for pk in documento.lines.values_list("product_id", flat=True) if pk]
    context["product_costs_json"] = {
        str(pk): str(costo)
        for pk, costo in Product.objects.filter(pk__in=ids, purchase_price__gt=0).values_list("pk", "purchase_price")
    }
    return context

def save_document_lines(document, formset, fk_field):
    lines = formset.save(commit=False)
    for line in lines:
        setattr(line, fk_field, document)
        # i moduli che non inviano posizione e tipo (fatture, ordini fornitore)
        # ricadono sui valori di default
        if line.position is None:
            line.position = 0
        if not line.line_type:
            line.line_type = LINE_ARTICLE
        line.save()
    eliminati = {obj.pk for obj in formset.deleted_objects}
    for line in formset.deleted_objects:
        line.delete()
    # L'ordine delle righe lo decide l'utente con le frecce su/giù: se il modulo
    # invia la posizione (preventivi, ordini, modelli) si rispetta quella,
    # altrimenti si ricade sull'ordine di creazione (comportamento storico).
    if "position" in getattr(formset.form, "base_fields", {}):
        righe = [f.instance for f in formset.forms if f.instance.pk and f.instance.pk not in eliminati]
        righe.sort(key=lambda riga: (riga.position or 0, riga.pk))
    else:
        righe = list(document.lines.order_by("pk"))
    for position, line in enumerate(righe, start=1):
        if line.position != position:
            line.position = position
            line.save(update_fields=["position"])
    services.espandi_kit(document, fk_field)
    if hasattr(document, "recalculate"):
        document.recalculate()


# --------------------------------------------------------------- preventivi
class QuoteListView(RoleRequiredMixin, ListView):
    allowed_roles = QUOTE_ROLES
    model = Quote
    template_name = "sales/quote_list.html"
    context_object_name = "quotes"
    paginate_by = 25

    def get_queryset(self):
        queryset = Quote.objects.select_related("customer", "created_by").order_by("-date", "-pk")
        status = self.request.GET.get("stato", "")
        if status == "aperti":
            queryset = queryset.filter(status__in=[Quote.STATUS_DRAFT, Quote.STATUS_SENT])
        elif status:
            queryset = queryset.filter(status=status)
        customer_id = self.request.GET.get("cliente", "")
        if customer_id:
            queryset = queryset.filter(customer_id=customer_id)
        # «miei=1» mostra solo i preventivi dell'utente collegato
        if self.request.GET.get("miei") == "1" and self.request.user.is_authenticated:
            queryset = queryset.filter(created_by=self.request.user)
        creator_id = self.request.GET.get("utente", "")
        if creator_id:
            queryset = queryset.filter(created_by_id=creator_id)
        search = self.request.GET.get("q", "").strip()
        if search:
            queryset = queryset.filter(Q(number__icontains=search) | Q(customer__name__icontains=search) | Q(reference__icontains=search))
        return queryset

    def get_context_data(self, **kwargs):
        from django.contrib.auth import get_user_model

        context = super().get_context_data(**kwargs)
        context["page_title"] = "Preventivi"
        context["statuses"] = Quote.STATUS_CHOICES
        context["customers"] = Contact.objects.filter(is_customer=True, active=True).order_by("name")
        context["status"] = self.request.GET.get("stato", "")
        context["customer_id"] = self.request.GET.get("cliente", "")
        context["search"] = self.request.GET.get("q", "")
        context["creator_id"] = self.request.GET.get("utente", "")
        context["only_mine"] = self.request.GET.get("miei") == "1"
        # solo chi ha davvero preventivi, per non riempire il menu di voci inutili
        context["creators"] = (
            get_user_model()
            .objects.filter(quotes__isnull=False)
            .distinct()
            .order_by("username")
        )
        return context


class QuoteDetailView(RoleRequiredMixin, DetailView):
    allowed_roles = QUOTE_ROLES
    model = Quote
    template_name = "sales/quote_detail.html"
    context_object_name = "quote"

    def get_queryset(self):
        return Quote.objects.select_related("customer", "payment_term", "created_by")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from apps.core.pdf import pdf_available

        context["page_title"] = f"Preventivo {self.object.number}"
        context["lines"] = self.object.lines.select_related("product", "uom", "vat_rate")
        context["line_groups"] = group_lines_by_section(context["lines"])
        context["vat_rows"] = self.object.vat_breakdown()
        context["generated_order"] = self.object.generated_order
        context["email_configured"] = emailing.email_configured()
        context["pdf_available"] = pdf_available()

        from apps.core.email_templates import contenuto
        from apps.core.models import CompanySettings, EmailTemplate
        from apps.core.utils import format_money

        contesto = {
            "cliente": self.object.customer.name,
            "numero": self.object.number,
            "data": f"{self.object.date:%d/%m/%Y}",
            "totale": format_money(self.object.grand_total),
            "azienda": CompanySettings.load().name,
            "cantiere": str(self.object.job) if self.object.job_id else "",
            "riferimento": self.object.reference or "",
            "validita": (
                f"Il preventivo è valido fino al {self.object.valid_until:%d/%m/%Y}.\n"
                if self.object.valid_until
                else ""
            ),
            "termini": (
                f"Condizioni generali di vendita: {self.request.build_absolute_uri(reverse('core:termini'))}\n"
            ),
        }
        context["email_subject"], context["email_body"] = contenuto(EmailTemplate.KIND_QUOTE, contesto)
        return context


class QuoteCreateView(RoleRequiredMixin, CreateView):
    allowed_roles = QUOTE_ROLES
    model = Quote
    form_class = QuoteForm
    template_name = "sales/quote_form.html"

    def get_success_url(self):
        return reverse_lazy("sales:quote_detail", args=[self.object.pk])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = "Nuovo preventivo"
        if "line_formset" not in context:
            context["line_formset"] = QuoteLineFormSet(prefix="lines", queryset=QuoteLine.objects.none())
        context["cancel_url"] = reverse("sales:quote_list")
        context["quote_templates"] = QuoteTemplate.objects.filter(is_active=True)
        return with_vat_rates(context)

    def post(self, request, *args, **kwargs):
        self.object = None
        form = self.get_form()
        formset = QuoteLineFormSet(request.POST, prefix="lines", queryset=QuoteLine.objects.none())
        if form.is_valid() and formset.is_valid():
            self.object = form.save(commit=False)
            self.object.created_by = request.user
            self.object.save()
            save_document_lines(self.object, formset, "quote")
            self.object.recalculate()
            messages.success(request, f"Preventivo «{self.object.number}» creato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


class QuoteUpdateView(ConflictAwareUpdateView, RoleRequiredMixin, UpdateView):
    allowed_roles = QUOTE_ROLES
    model = Quote
    form_class = QuoteForm
    template_name = "sales/quote_form.html"

    def get_success_url(self):
        return reverse_lazy("sales:quote_detail", args=[self.object.pk])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = f"Modifica preventivo {self.object.number}"
        if "line_formset" not in context:
            context["line_formset"] = QuoteLineFormSet(prefix="lines", queryset=self.object.lines.select_related("product"))
        context["cancel_url"] = reverse("sales:quote_detail", args=[self.object.pk])
        return with_vat_rates(context, self.object)

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.get_form()
        formset = QuoteLineFormSet(request.POST, prefix="lines", queryset=self.object.lines.select_related("product"))
        if form.is_valid() and formset.is_valid():
            self.object = form.save()
            save_document_lines(self.object, formset, "quote")
            self.object.recalculate()
            messages.success(request, f"Preventivo «{self.object.number}» aggiornato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


class QuotePrintView(RoleRequiredMixin, DetailView):
    allowed_roles = QUOTE_ROLES
    model = Quote
    template_name = "print/document.html"
    context_object_name = "quote"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # import locale: printing.py usa funzioni di questo modulo, quindi
        # importarlo in cima creerebbe un ciclo
        from .printing import quote_print_context

        context.update(quote_print_context(self.object))
        return context


def _pdf_download_response(request, document, numero, render, stampa_url):
    """File PDF generato dal gestionale (niente scritte del browser)."""
    from io import BytesIO

    from django.http import FileResponse

    pdf = render(document)
    if pdf is None:
        messages.warning(request, "PDF non disponibile su questo sistema: usa «Stampa / Salva PDF» del browser.")
        return redirect(stampa_url, pk=document.pk)
    return FileResponse(BytesIO(pdf), as_attachment=True, filename=f"{numero}.pdf", content_type="application/pdf")


@role_required(*QUOTE_ROLES)
def quote_pdf_download(request, pk):
    """Scarica il PDF del preventivo generato dal gestionale."""
    from .pdf import render_quote_pdf

    quote = get_object_or_404(Quote, pk=pk)
    return _pdf_download_response(request, quote, quote.number, render_quote_pdf, "sales:quote_print")


@role_required(*QUOTE_ROLES)
def quote_send(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    if quote.status == Quote.STATUS_DRAFT:
        quote.status = Quote.STATUS_SENT
        quote.save(update_fields=["status"])
        messages.success(request, f"Preventivo {quote.number} segnato come inviato.")
    else:
        messages.warning(request, "Il preventivo non è in bozza: stato non modificato.")
    return redirect("sales:quote_detail", pk=quote.pk)


@role_required(*QUOTE_ROLES)
def quote_email(request, pk):
    """Accoda il preventivo per l'invio email (con PDF in allegato se disponibile).

    L'invio vero avviene in background ogni 2 minuti: la pagina risponde
    subito anche se il server di posta è lento.
    """
    from apps.core.mailing import accoda_email
    from apps.core.models import EmailInCoda

    quote = get_object_or_404(Quote.objects.select_related("customer"), pk=pk)
    if request.method == "POST":
        to_email = (request.POST.get("to") or quote.customer.email or "").strip()
        subject = (request.POST.get("subject") or f"Preventivo {quote.number}").strip()
        message_body = (request.POST.get("message") or "").strip()

        if not to_email:
            messages.error(request, "Indica l'indirizzo email del cliente.")
        elif not emailing.email_configured():
            messages.error(request, "Invio email non configurato: imposta le variabili EMAIL_* nel file .env.")
        else:
            try:
                pdf = emailing.render_quote_pdf(quote)
                accoda_email(
                    to_email=to_email,
                    subject=subject,
                    message=message_body,
                    attachment=pdf,
                    attachment_name=f"Preventivo_{quote.number}.pdf",
                    descrizione=f"Preventivo {quote.number} a {to_email}",
                    modello=EmailInCoda.MODELLO_PREVENTIVO,
                    oggetto_id=quote.pk,
                )
            except Exception as exc:
                messages.error(request, f"Accodamento non riuscito: {exc}")
            else:
                extra = "" if pdf else " (senza allegato PDF: non disponibile su questo sistema)"
                messages.success(
                    request,
                    f"Preventivo {quote.number} accodato per {to_email}{extra}: verrà inviato entro pochi minuti.",
                )
    return redirect("sales:quote_detail", pk=quote.pk)


@role_required(*QUOTE_ROLES)
def quote_accept(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    if quote.status in {Quote.STATUS_DRAFT, Quote.STATUS_SENT}:
        quote.status = Quote.STATUS_ACCEPTED
        quote.save(update_fields=["status"])
        messages.success(request, f"Preventivo {quote.number} accettato dal cliente.")
    else:
        messages.warning(request, "Il preventivo non può essere accettato nello stato attuale.")
    return redirect("sales:quote_detail", pk=quote.pk)


@role_required(*QUOTE_ROLES)
def quote_reject(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    if quote.status in {Quote.STATUS_DRAFT, Quote.STATUS_SENT, Quote.STATUS_ACCEPTED}:
        quote.status = Quote.STATUS_REJECTED
        quote.save(update_fields=["status"])
        messages.info(request, f"Preventivo {quote.number} rifiutato.")
    else:
        messages.warning(request, "Il preventivo non può essere rifiutato nello stato attuale.")
    return redirect("sales:quote_detail", pk=quote.pk)


@role_required(*QUOTE_ROLES)
def quote_convert(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    try:
        order = services.convert_quote_to_order(quote, user=request.user)
    except ValidationError as exc:
        messages.error(request, "; ".join(exc.messages))
        return redirect("sales:quote_detail", pk=quote.pk)
    messages.success(request, f"Creato l'ordine cliente {order.number} dal preventivo {quote.number}.")
    return redirect("sales:order_detail", pk=order.pk)


@role_required(*QUOTE_ROLES)
def quote_duplicate(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    new = services.duplicate_quote(quote, user=request.user)
    messages.success(request, f"Preventivo duplicato: {new.number}.")
    return redirect("sales:quote_update", pk=new.pk)


@role_required(*QUOTE_ROLES)
def quote_delete(request, pk):
    quote = get_object_or_404(Quote, pk=pk)
    if request.method == "POST":
        if quote.generated_orders.exists():
            messages.error(request, "Il preventivo è collegato a un ordine cliente: non può essere eliminato.")
        else:
            number = quote.number
            quote.delete()
            messages.success(request, f"Preventivo {number} eliminato.")
    return redirect("sales:quote_list")


# ----------------------------------------------------------- ordini cliente
class SalesOrderListView(RoleRequiredMixin, ListView):
    allowed_roles = ORDER_VIEW_ROLES
    model = SalesOrder
    template_name = "sales/order_list.html"
    context_object_name = "orders"
    paginate_by = 25

    def get_queryset(self):
        queryset = SalesOrder.objects.select_related("customer").order_by("-date", "-pk")
        status = self.request.GET.get("stato", "")
        if status == "aperti":
            queryset = queryset.filter(status__in=[SalesOrder.STATUS_DRAFT, SalesOrder.STATUS_CONFIRMED])
        elif status:
            queryset = queryset.filter(status=status)
        customer_id = self.request.GET.get("cliente", "")
        if customer_id:
            queryset = queryset.filter(customer_id=customer_id)
        search = self.request.GET.get("q", "").strip()
        if search:
            queryset = queryset.filter(Q(number__icontains=search) | Q(customer__name__icontains=search) | Q(reference__icontains=search))
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = "Ordini cliente"
        context["statuses"] = SalesOrder.STATUS_CHOICES
        context["customers"] = Contact.objects.filter(is_customer=True, active=True).order_by("name")
        context["status"] = self.request.GET.get("stato", "")
        context["customer_id"] = self.request.GET.get("cliente", "")
        context["search"] = self.request.GET.get("q", "")
        return context


class SalesOrderDetailView(RoleRequiredMixin, DetailView):
    allowed_roles = ORDER_VIEW_ROLES
    model = SalesOrder
    template_name = "sales/order_detail.html"
    context_object_name = "order"

    def get_queryset(self):
        return SalesOrder.objects.select_related("customer", "payment_term", "created_by", "source_quote")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = f"Ordine {self.object.number}"
        context["lines"] = self.object.lines.select_related("product", "uom", "vat_rate")
        context["line_groups"] = group_lines_by_section(context["lines"])
        context["vat_rows"] = self.object.vat_breakdown()
        context["purchase_orders"] = self.object.purchase_orders.select_related("supplier").order_by("pk")
        from apps.billing.services import order_billing_summary
        from apps.core.pdf import pdf_available

        context["billing"] = order_billing_summary(self.object)
        context["pdf_available"] = pdf_available()
        return context


class SalesOrderCreateView(RoleRequiredMixin, CreateView):
    allowed_roles = ORDER_EDIT_ROLES
    model = SalesOrder
    form_class = SalesOrderForm
    template_name = "sales/order_form.html"

    def get_success_url(self):
        return reverse_lazy("sales:order_detail", args=[self.object.pk])

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = "Nuovo ordine cliente"
        if "line_formset" not in context:
            context["line_formset"] = SalesOrderLineFormSet(prefix="lines", queryset=SalesOrderLine.objects.none())
        context["cancel_url"] = reverse("sales:order_list")
        return with_vat_rates(context)

    def post(self, request, *args, **kwargs):
        self.object = None
        form = self.get_form()
        formset = SalesOrderLineFormSet(request.POST, prefix="lines", queryset=SalesOrderLine.objects.none())
        if form.is_valid() and formset.is_valid():
            self.object = form.save(commit=False)
            self.object.created_by = request.user
            self.object.save()
            save_document_lines(self.object, formset, "order")
            self.object.recalculate()
            messages.success(request, f"Ordine «{self.object.number}» creato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


class SalesOrderUpdateView(ConflictAwareUpdateView, RoleRequiredMixin, UpdateView):
    allowed_roles = ORDER_EDIT_ROLES
    model = SalesOrder
    form_class = SalesOrderForm
    template_name = "sales/order_form.html"

    def get_success_url(self):
        return reverse_lazy("sales:order_detail", args=[self.object.pk])

    def dispatch(self, request, *args, **kwargs):
        self.object = self.get_object()
        if not self.object.is_editable:
            messages.warning(request, "L'ordine non è modificabile nello stato attuale.")
            return redirect("sales:order_detail", pk=self.object.pk)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = f"Modifica ordine {self.object.number}"
        if "line_formset" not in context:
            context["line_formset"] = SalesOrderLineFormSet(prefix="lines", queryset=self.object.lines.select_related("product"))
        context["cancel_url"] = reverse("sales:order_detail", args=[self.object.pk])
        return with_vat_rates(context, self.object)

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.get_form()
        formset = SalesOrderLineFormSet(request.POST, prefix="lines", queryset=self.object.lines.select_related("product"))
        if form.is_valid() and formset.is_valid():
            self.object = form.save()
            save_document_lines(self.object, formset, "order")
            self.object.recalculate()
            messages.success(request, f"Ordine «{self.object.number}» aggiornato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


class SalesOrderPrintView(RoleRequiredMixin, DetailView):
    allowed_roles = ORDER_VIEW_ROLES
    model = SalesOrder
    template_name = "print/document.html"
    context_object_name = "order"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        doc = self.object
        # import locale: printing.py usa funzioni di questo modulo, quindi
        # importarlo in cima creerebbe un ciclo
        from .printing import order_print_context

        context.update(order_print_context(doc))
        return context


@role_required(*ORDER_VIEW_ROLES)
def order_pdf_download(request, pk):
    """Scarica il PDF della conferma d'ordine generato dal gestionale."""
    from .pdf import render_order_pdf

    order = get_object_or_404(SalesOrder, pk=pk)
    return _pdf_download_response(request, order, order.number, render_order_pdf, "sales:order_print")


@role_required(*ORDER_EDIT_ROLES)
def order_confirm(request, pk):
    order = get_object_or_404(SalesOrder, pk=pk)
    replenish = request.method == "POST" and request.POST.get("replenish") == "1"
    try:
        result = services.confirm_sales_order(order, user=request.user, replenish=replenish)
    except ValidationError as exc:
        messages.error(request, "; ".join(exc.messages))
        return redirect("sales:order_detail", pk=order.pk)

    extra = " Sono stati riordinati anche gli articoli sotto scorta minima." if replenish else ""
    messages.success(request, f"Ordine {order.number} confermato.{extra}")
    purchase_orders = result["purchase_orders"]
    if purchase_orders:
        numbers = ", ".join(po.number for po in purchase_orders)
        messages.info(request, f"Creati {len(purchase_orders)} ordini fornitore per le carenze di magazzino: {numbers}.")
    for product, shortage in result["without_supplier"]:
        messages.warning(
            request,
            f"«{product.name}»: mancano {shortage} e l'articolo non ha un fornitore abituale. "
            "Imposta il fornitore nell'articolo o crea l'ordine fornitore manualmente.",
        )
    return redirect("sales:order_detail", pk=order.pk)


@role_required(*ORDER_DELIVER_ROLES)
def order_deliver(request, pk):
    order = get_object_or_404(SalesOrder, pk=pk)
    if order.status == SalesOrder.STATUS_DRAFT:
        messages.warning(request, "Conferma l'ordine prima di registrare la consegna.")
        return redirect("sales:order_detail", pk=order.pk)
    try:
        errors = services.deliver_sales_order(order, user=request.user)
    except ValidationError as exc:
        messages.error(request, "; ".join(exc.messages))
        return redirect("sales:order_detail", pk=order.pk)

    if errors:
        for error in errors:
            messages.error(request, error)
        messages.warning(request, "Consegna registrata parzialmente: alcune righe non avevano giacenza sufficiente.")
    else:
        messages.success(request, f"Consegna dell'ordine {order.number} registrata e magazzino scaricato.")
    return redirect("sales:order_detail", pk=order.pk)


@role_required(*ORDER_EDIT_ROLES)
def order_cancel(request, pk):
    order = get_object_or_404(SalesOrder, pk=pk)
    if request.method == "POST":
        try:
            services.cancel_sales_order(order, user=request.user)
            messages.info(request, f"Ordine {order.number} annullato.")
        except ValidationError as exc:
            messages.error(request, "; ".join(exc.messages))
    return redirect("sales:order_detail", pk=order.pk)


@role_required(*ORDER_EDIT_ROLES)
def order_delete(request, pk):
    order = get_object_or_404(SalesOrder, pk=pk)
    if request.method == "POST":
        if order.status != SalesOrder.STATUS_DRAFT:
            messages.error(request, "Solo gli ordini in bozza possono essere eliminati.")
        else:
            number = order.number
            order.delete()
            messages.success(request, f"Ordine {number} eliminato.")
            return redirect("sales:order_list")
    return redirect("sales:order_detail", pk=order.pk)


# --------------------------------------------------- modelli di preventivo
class QuoteTemplateListView(RoleRequiredMixin, ListView):
    allowed_roles = QUOTE_ROLES
    model = QuoteTemplate
    template_name = "sales/quotetemplate_list.html"
    context_object_name = "templates"
    paginate_by = 50

    def get_queryset(self):
        return QuoteTemplate.objects.annotate(line_count=Count("lines")).order_by("sort_order", "name")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = "Modelli di preventivo"
        return context


class QuoteTemplateCreateView(RoleRequiredMixin, CreateView):
    allowed_roles = QUOTE_ROLES
    model = QuoteTemplate
    form_class = QuoteTemplateForm
    template_name = "sales/quotetemplate_form.html"
    success_url = reverse_lazy("sales:quote_template_list")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = "Nuovo modello di preventivo"
        if "line_formset" not in context:
            context["line_formset"] = QuoteTemplateLineFormSet(prefix="lines", queryset=QuoteTemplateLine.objects.none())
        context["cancel_url"] = reverse("sales:quote_template_list")
        return with_vat_rates(context)

    def post(self, request, *args, **kwargs):
        self.object = None
        form = self.get_form()
        formset = QuoteTemplateLineFormSet(request.POST, prefix="lines", queryset=QuoteTemplateLine.objects.none())
        if form.is_valid() and formset.is_valid():
            self.object = form.save(commit=False)
            self.object.created_by = request.user
            self.object.save()
            save_document_lines(self.object, formset, "template")
            messages.success(request, f"Modello «{self.object.name}» creato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


class QuoteTemplateUpdateView(ConflictAwareUpdateView, RoleRequiredMixin, UpdateView):
    allowed_roles = QUOTE_ROLES
    model = QuoteTemplate
    form_class = QuoteTemplateForm
    template_name = "sales/quotetemplate_form.html"
    success_url = reverse_lazy("sales:quote_template_list")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_title"] = f"Modifica modello: {self.object.name}"
        if "line_formset" not in context:
            context["line_formset"] = QuoteTemplateLineFormSet(prefix="lines", queryset=self.object.lines.select_related("product"))
        context["cancel_url"] = reverse("sales:quote_template_list")
        return with_vat_rates(context, self.object)

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.get_form()
        formset = QuoteTemplateLineFormSet(request.POST, prefix="lines", queryset=self.object.lines.select_related("product"))
        if form.is_valid() and formset.is_valid():
            self.object = form.save()
            save_document_lines(self.object, formset, "template")
            messages.success(request, f"Modello «{self.object.name}» aggiornato.")
            return redirect(self.get_success_url())
        return self.render_to_response(self.get_context_data(form=form, line_formset=formset))


@role_required(*QUOTE_ROLES)
def quote_template_delete(request, pk):
    template = get_object_or_404(QuoteTemplate, pk=pk)
    if request.method == "POST":
        name = template.name
        template.delete()
        messages.success(request, f"Modello «{name}» eliminato.")
    return redirect("sales:quote_template_list")


@role_required(*QUOTE_ROLES)
def quote_template_data(request, pk):
    """Dati del modello per il riempimento del preventivo (JSON)."""
    template = get_object_or_404(QuoteTemplate, pk=pk)
    lines = []
    for line in template.lines.select_related("product", "uom", "vat_rate"):
        lines.append(
            {
                "product": line.product_id,
                "product_label": f"{line.product.code} – {line.product.name}" if line.product_id else "",
                "description": line.description or (line.product.name if line.product_id else ""),
                "line_type": line.line_type,
                "qty": plain_number(line.qty),
                "uom": line.uom_id,
                "unit_price": plain_number(line.unit_price),
                "discount_pct": plain_number(line.discount_pct),
                "vat_rate": line.vat_rate_id,
            }
        )
    return JsonResponse(
        {
            "template": {
                "id": template.pk,
                "name": template.name,
                "payment_term": template.payment_term_id,
                "terms_text": template.terms_text,
                "notes": template.notes,
            },
            "lines": lines,
        }
    )


# ------------------------------------------------------------ esportazione
EXPORT_COLUMNS = [
    ("Data", 11),
    ("Documento", 14),
    ("Cliente", 30),
    ("Venditore", 16),
    ("Codice", 14),
    ("Articolo / descrizione", 42),
    ("Q.tà", 10),
    ("U.d.M.", 8),
    ("Prezzo unitario", 14),
    ("Sconto %", 10),
    ("Imponibile", 13),
    ("Costo", 13),
    ("Provvigione", 13),
    ("Margine", 13),
    ("Margine %", 11),
]


@role_required(*QUOTE_ROLES)
def export_sales_excel(request):
    """Scarica in Excel il dettaglio delle vendite del periodo scelto.

    Una riga per articolo venduto, con costo, provvigione e margine, così il
    file si può aprire, filtrare e pivotare direttamente.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    period = request.GET.get("periodo", analytics.PERIOD_YEAR)
    if period not in {value for value, _label in analytics.PERIOD_CHOICES}:
        period = analytics.PERIOD_YEAR

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vendite"

    intestazione = Font(bold=True, color="FFFFFF")
    sfondo = PatternFill("solid", fgColor="2563EB")
    for indice, (titolo, larghezza) in enumerate(EXPORT_COLUMNS, start=1):
        cella = sheet.cell(row=1, column=indice, value=titolo)
        cella.font = intestazione
        cella.fill = sfondo
        cella.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(indice)].width = larghezza
    sheet.freeze_panes = "A2"

    totali = {"imponibile": Decimal("0"), "costo": Decimal("0"), "provvigione": Decimal("0"), "margine": Decimal("0")}
    riga = 2
    for line, line_revenue, row_cost, share in analytics.lines_with_shares(period):
        order = line.order
        margine = line_revenue - row_cost - share
        valori = [
            timezone.localtime(order.delivered_at).date() if order.delivered_at else order.date,
            order.number,
            order.customer.name,
            (order.created_by.get_full_name() or order.created_by.username) if order.created_by else "",
            line.product.code if line.product_id else "",
            line.description or (line.product.name if line.product_id else ""),
            float(line.qty or 0),
            line.uom.code if line.uom_id else "",
            float(line.unit_price or 0),
            float(line.discount_pct or 0),
            float(line_revenue),
            float(row_cost),
            float(share),
            float(margine),
            float(round(margine / line_revenue * 100, 2)) if line_revenue else 0.0,
        ]
        for indice, valore in enumerate(valori, start=1):
            cella = sheet.cell(row=riga, column=indice, value=valore)
            if indice == 1:
                cella.number_format = "DD/MM/YYYY"
            elif indice in (7, 9, 10, 11, 12, 13, 14, 15):
                cella.number_format = "#,##0.00"
        totali["imponibile"] += line_revenue
        totali["costo"] += row_cost
        totali["provvigione"] += share
        totali["margine"] += margine
        riga += 1

    # riga dei totali, comoda per verificare che i conti tornino
    if riga > 2:
        sheet.cell(row=riga, column=6, value="TOTALE").font = Font(bold=True)
        for colonna, chiave in ((11, "imponibile"), (12, "costo"), (13, "provvigione"), (14, "margine")):
            cella = sheet.cell(row=riga, column=colonna, value=float(totali[chiave]))
            cella.font = Font(bold=True)
            cella.number_format = "#,##0.00"

    nome = f"vendite_{period}_{timezone.localdate():%Y-%m-%d}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{nome}"'
    workbook.save(response)
    return response
