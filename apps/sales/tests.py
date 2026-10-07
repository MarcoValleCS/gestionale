"""Test del flusso di vendita completo e delle pagine principali."""
import re
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from datetime import timedelta

from apps.catalog.models import Category, Product
from apps.contacts.models import Contact
from apps.core.models import CompanySettings, Tag, UnitOfMeasure, VatRate
from apps.inventory.models import StockLevel, Warehouse
from apps.inventory.services import register_movement
from apps.jobs.models import Job
from apps.purchasing.models import PurchaseOrder, SupplierPriceList, PriceListItem
from apps.purchasing.services import apply_pricelist_adjustment

from . import analytics, services
from .models import Quote, QuoteLine, QuoteTemplate, SalesOrder, group_lines_by_section

User = get_user_model()


class FlowTestBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_superuser("admin", "admin@example.com", "password123!")
        cls.uom = UnitOfMeasure.objects.get(code="PZ")
        cls.vat22 = VatRate.objects.get(code="22")
        cls.vat10 = VatRate.objects.get(code="10")
        cls.warehouse = Warehouse.get_default()

        cls.customer = Contact.objects.create(name="Cliente Test S.r.l.", is_customer=True, is_supplier=False, city="Milano")
        cls.supplier = Contact.objects.create(name="Fornitore Test S.p.A.", is_customer=False, is_supplier=True, city="Torino")

        cls.product = Product.objects.create(
            name="Bullone M8",
            uom=cls.uom,
            sale_price=Decimal("10.00"),
            sale_vat=cls.vat22,
            purchase_price=Decimal("5.00"),
            purchase_vat=cls.vat22,
            main_supplier=cls.supplier,
            min_stock=Decimal("20"),
        )

    def login(self):
        self.client.force_login(self.user)


class FullFlowTest(FlowTestBase):
    def test_preventivo_ordine_acquisti_consegna(self):
        self.login()

        # 1. Creazione preventivo tramite form web
        response = self.client.post(
            reverse("sales:quote_create"),
            {
                "customer": self.customer.pk,
                "date": "2026-01-10",
                "valid_until": "2026-02-10",
                "payment_term": "",
                "reference": "Rif cliente 1",
                "terms_text": "",
                "notes": "",
                "lines-TOTAL_FORMS": "1",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-product": self.product.pk,
                "lines-0-description": "",
                "lines-0-qty": "10",
                "lines-0-uom": self.uom.pk,
                "lines-0-unit_price": "10.00",
                "lines-0-discount_pct": "0",
                "lines-0-vat_rate": self.vat22.pk,
            },
        )
        self.assertEqual(response.status_code, 302, getattr(response, "context", None))
        quote = Quote.objects.get()
        self.assertTrue(quote.number.startswith("PRE-"))
        self.assertEqual(quote.subtotal, Decimal("100.00"))
        self.assertEqual(quote.vat_total, Decimal("22.00"))
        self.assertEqual(quote.grand_total, Decimal("122.00"))
        line = quote.lines.get()
        self.assertEqual(line.description, "Bullone M8")  # riempita dall'articolo

        # 2. Invio e accettazione
        self.client.post(reverse("sales:quote_send", args=[quote.pk]))
        quote.refresh_from_db()
        self.assertEqual(quote.status, Quote.STATUS_SENT)

        self.client.post(reverse("sales:quote_accept", args=[quote.pk]))
        quote.refresh_from_db()
        self.assertEqual(quote.status, Quote.STATUS_ACCEPTED)

        # 3. Conversione in ordine cliente
        self.client.post(reverse("sales:quote_convert", args=[quote.pk]))
        quote.refresh_from_db()
        self.assertEqual(quote.status, Quote.STATUS_CONVERTED)
        order = SalesOrder.objects.get()
        self.assertTrue(order.number.startswith("OC-"))
        self.assertEqual(order.customer, self.customer)
        self.assertEqual(order.lines.count(), 1)

        # 4. Conferma ordine → nessuna giacenza → ordine fornitore automatico
        self.client.post(reverse("sales:order_confirm", args=[order.pk]))
        order.refresh_from_db()
        self.assertEqual(order.status, SalesOrder.STATUS_CONFIRMED)
        po = PurchaseOrder.objects.get()
        self.assertEqual(po.supplier, self.supplier)
        self.assertEqual(po.source_sales_order, order)
        po_line = po.lines.get()
        self.assertEqual(po_line.qty, Decimal("10"))
        self.assertEqual(po_line.unit_price, Decimal("5.00"))

        # 5. Ricezione merce → carico magazzino
        self.client.post(reverse("purchasing:po_receive", args=[po.pk]), {"receive_all": "1"})
        po.refresh_from_db()
        self.assertEqual(po.status, PurchaseOrder.STATUS_RECEIVED)
        level = StockLevel.objects.get(product=self.product, warehouse=self.warehouse)
        self.assertEqual(level.quantity, Decimal("10"))

        # 6. Consegna → scarico magazzino
        self.client.post(reverse("sales:order_deliver", args=[order.pk]))
        order.refresh_from_db()
        self.assertEqual(order.status, SalesOrder.STATUS_DELIVERED)
        level.refresh_from_db()
        self.assertEqual(level.quantity, Decimal("0"))

    def test_conferma_senza_fornitore_non_crea_ordine(self):
        product_no_supplier = Product.objects.create(
            name="Articolo senza fornitore",
            uom=self.uom,
            sale_price=Decimal("3.00"),
            sale_vat=self.vat22,
            purchase_price=Decimal("1.00"),
            purchase_vat=self.vat22,
        )
        order = SalesOrder.objects.create(customer=self.customer)
        order.lines.create(product=product_no_supplier, description="x", qty=Decimal("2"), uom=self.uom, unit_price=Decimal("3"), vat_rate=self.vat22)
        order.recalculate()

        result = services.confirm_sales_order(order, user=self.user)
        self.assertEqual(len(result["purchase_orders"]), 0)
        self.assertEqual(len(result["without_supplier"]), 1)
        order.refresh_from_db()
        self.assertEqual(order.status, SalesOrder.STATUS_CONFIRMED)

    def test_consegna_bloccata_senza_giacenza(self):
        order = SalesOrder.objects.create(customer=self.customer)
        order.lines.create(product=self.product, description="x", qty=Decimal("5"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        order.recalculate()
        services.confirm_sales_order(order, user=self.user)

        errors = services.deliver_sales_order(order, user=self.user)
        self.assertEqual(len(errors), 1)
        order.refresh_from_db()
        self.assertEqual(order.status, SalesOrder.STATUS_CONFIRMED)
        self.assertFalse(StockLevel.objects.filter(product=self.product).exists())

    def test_duplicazione_preventivo(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(product=self.product, description="x", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        quote.recalculate()
        new_quote = services.duplicate_quote(quote, user=self.user)
        self.assertNotEqual(new_quote.pk, quote.pk)
        self.assertEqual(new_quote.status, Quote.STATUS_DRAFT)
        self.assertEqual(new_quote.grand_total, quote.grand_total)
        self.assertEqual(new_quote.lines.count(), 1)

    def test_conferma_con_riordino_scorte(self):
        # min_stock = 20, ordine 10, giacenza 0 → il riordino porta a 20
        order = SalesOrder.objects.create(customer=self.customer)
        order.lines.create(product=self.product, description="x", qty=Decimal("10"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        order.recalculate()
        result = services.confirm_sales_order(order, user=self.user, replenish=True)
        po = result["purchase_orders"][0]
        self.assertEqual(po.lines.get().qty, Decimal("20"))

    def test_ordine_fornitore_usa_giorni_consegna(self):
        self.product.supplier_lead_days = 5
        self.product.save(update_fields=["supplier_lead_days"])
        order = SalesOrder.objects.create(customer=self.customer)
        order.lines.create(product=self.product, description="x", qty=Decimal("3"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        order.recalculate()
        result = services.confirm_sales_order(order, user=self.user)
        po = result["purchase_orders"][0]
        self.assertEqual(po.expected_date, timezone.localdate() + timedelta(days=5))

    def test_raggruppamento_per_sezioni(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(section="Bagno 1", product=self.product, description="a", qty=Decimal("2"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        quote.lines.create(section="Bagno 1", product=self.product, description="b", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("5"), vat_rate=self.vat22)
        quote.lines.create(section="Bagno 2", product=self.product, description="c", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("100"), vat_rate=self.vat22)
        groups = group_lines_by_section(quote.lines.order_by("position", "pk"))
        self.assertEqual(len(groups), 2)
        self.assertEqual(groups[0]["section"], "Bagno 1")
        self.assertEqual(groups[0]["subtotal"], Decimal("25.00"))
        self.assertEqual(groups[1]["subtotal"], Decimal("100.00"))
        self.assertEqual(groups[1]["lines"][0].row_number, 3)


class AnalyticsTest(FlowTestBase):
    def _delivered_order(self, qty="10", price="10.00", unit_cost="4.00"):
        order = SalesOrder.objects.create(
            customer=self.customer,
            status=SalesOrder.STATUS_DELIVERED,
            delivered_at=timezone.now(),
        )
        order.lines.create(
            product=self.product,
            description=self.product.name,
            qty=Decimal(qty),
            uom=self.uom,
            unit_price=Decimal(price),
            vat_rate=self.vat22,
            qty_delivered=Decimal(qty),
            unit_cost=Decimal(unit_cost) if unit_cost is not None else None,
        )
        order.recalculate()
        return order

    def test_fatturato_e_margine(self):
        self._delivered_order()

        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["revenue"], Decimal("100.00"))
        self.assertEqual(summary["cost"], Decimal("40.00"))
        self.assertEqual(summary["margin"], Decimal("60.00"))
        self.assertEqual(summary["margin_pct"], Decimal("60.00"))
        self.assertEqual(summary["orders"], 1)

        by_product = analytics.by_product(analytics.PERIOD_YEAR)
        self.assertEqual(len(by_product), 1)
        self.assertEqual(by_product[0]["label"], "Bullone M8")
        self.assertEqual(by_product[0]["margin"], Decimal("60.00"))

        by_supplier = analytics.by_supplier(analytics.PERIOD_YEAR)
        self.assertEqual(by_supplier[0]["label"], "Fornitore Test S.p.A.")
        self.assertEqual(by_supplier[0]["margin"], Decimal("60.00"))

    def test_margine_usa_prezzo_acquisto_se_costo_non_fissato(self):
        self._delivered_order(unit_cost=None)
        summary = analytics.summary(analytics.PERIOD_YEAR)
        # prezzo di acquisto dell'articolo = 5.00 → costo 50
        self.assertEqual(summary["cost"], Decimal("50.00"))
        self.assertEqual(summary["margin"], Decimal("50.00"))

    def test_serie_mensile(self):
        self._delivered_order()
        series = analytics.monthly_series(12)
        self.assertEqual(len(series["labels"]), 12)
        self.assertEqual(series["revenue"][-1], 100.0)
        self.assertEqual(series["margin"][-1], 60.0)
        self.assertEqual(series["revenue"][0], 0.0)

    def test_ordini_non_consegnati_esclusi(self):
        order = SalesOrder.objects.create(customer=self.customer, status=SalesOrder.STATUS_CONFIRMED)
        order.lines.create(product=self.product, description="x", qty=Decimal("5"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        order.recalculate()
        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["revenue"], Decimal("0"))

    def test_dashboard_con_statistiche(self):
        self._delivered_order()
        self.login()
        for period in ("anno", "mese", "12m"):
            with self.subTest(period=period):
                response = self.client.get(reverse("core:home") + f"?periodo={period}")
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Fatturato")
                self.assertContains(response, "Statistiche complete")

    def test_pagina_statistiche_con_dettagli(self):
        self._delivered_order()
        self.login()
        for period in ("anno", "mese", "12m"):
            with self.subTest(period=period):
                response = self.client.get(reverse("sales:statistics") + f"?periodo={period}")
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Marginalità per articolo")
                self.assertContains(response, "Fatturato per cliente")
                self.assertContains(response, "Fatturato per cantiere")

    def test_fatturato_per_cliente_e_per_cantiere(self):
        job = Job.objects.create(name="Piscina Rossi", customer=self.customer)
        order = self._delivered_order()
        order.job = job
        order.save(update_fields=["job"])

        by_customer = analytics.by_customer(analytics.PERIOD_YEAR)
        self.assertEqual(by_customer[0]["label"], "Cliente Test S.r.l.")
        self.assertEqual(by_customer[0]["revenue"], Decimal("100.00"))

        by_job = analytics.by_job(analytics.PERIOD_YEAR)
        self.assertEqual(len(by_job), 1)
        self.assertIn("Piscina Rossi", by_job[0]["label"])
        self.assertEqual(by_job[0]["margin"], Decimal("60.00"))

    def test_serie_mensile_finestra_e_periodo_precedente(self):
        self._delivered_order()
        series = analytics.monthly_series(3)
        self.assertEqual(len(series["labels"]), 3)
        self.assertEqual(len(series["rows"]), 3)
        self.assertEqual(series["revenue"][-1], 100.0)
        self.assertEqual(series["rows"][-1]["revenue"], Decimal("100.00"))

        previous = analytics.monthly_series(3, end_offset=1)
        self.assertEqual(len(previous["labels"]), 3)
        self.assertEqual(previous["revenue"][-1], 0.0)
        self.assertNotEqual(previous["labels"], series["labels"])


class QuoteTemplateTest(FlowTestBase):
    def test_crea_modello_tramite_vista(self):
        self.login()
        response = self.client.post(
            reverse("sales:quote_template_create"),
            {
                "name": "Bagno completo",
                "description": "Modello base bagno",
                "payment_term": "",
                "terms_text": "Validità 30 giorni. Posa inclusa.",
                "notes": "Note interne del modello",
                "sort_order": "1",
                "is_active": "on",
                "lines-TOTAL_FORMS": "1",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-product": self.product.pk,
                "lines-0-description": "",
                "lines-0-qty": "4",
                "lines-0-uom": self.uom.pk,
                "lines-0-unit_price": "10.00",
                "lines-0-discount_pct": "0",
                "lines-0-vat_rate": self.vat22.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        template = QuoteTemplate.objects.get()
        self.assertEqual(template.name, "Bagno completo")
        self.assertEqual(template.lines.count(), 1)
        self.assertEqual(template.lines.first().description, "Bullone M8")

    def test_dati_modello_json(self):
        template = QuoteTemplate.objects.create(name="Piscina 8x4", terms_text="Incluso scavo")
        template.lines.create(
            product=self.product,
            description="Scavo",
            qty=Decimal("2.500"),
            uom=self.uom,
            unit_price=Decimal("100.00"),
            discount_pct=Decimal("5"),
            vat_rate=self.vat22,
        )
        self.login()
        response = self.client.get(reverse("sales:quote_template_data", args=[template.pk]))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["template"]["name"], "Piscina 8x4")
        self.assertEqual(data["template"]["terms_text"], "Incluso scavo")
        self.assertEqual(len(data["lines"]), 1)
        line = data["lines"][0]
        self.assertEqual(line["qty"], "2.5")
        self.assertEqual(line["unit_price"], "100")
        self.assertEqual(line["discount_pct"], "5")
        self.assertEqual(line["product"], self.product.pk)
        self.assertEqual(line["vat_rate"], self.vat22.pk)

    def test_form_preventivo_mostra_i_modelli(self):
        QuoteTemplate.objects.create(name="Manutenzione annuale")
        self.login()
        response = self.client.get(reverse("sales:quote_create"))
        self.assertContains(response, "Manutenzione annuale")
        self.assertContains(response, "quote-template-select")
        self.assertContains(response, "quote_template.js")


class VatAndTotalsTest(FlowTestBase):
    def test_totali_con_iva_mista_e_sconto(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(product=self.product, description="a", qty=Decimal("2"), uom=self.uom, unit_price=Decimal("10"), discount_pct=Decimal("10"), vat_rate=self.vat22)
        quote.lines.create(product=self.product, description="b", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("100"), vat_rate=self.vat10)
        quote.recalculate()

        # Riga 1: 2*10*0.9 = 18 → IVA 22% = 3.96
        # Riga 2: 100 → IVA 10% = 10.00
        self.assertEqual(quote.subtotal, Decimal("118.00"))
        self.assertEqual(quote.vat_total, Decimal("13.96"))
        self.assertEqual(quote.grand_total, Decimal("131.96"))

        rows = quote.vat_breakdown()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["rate"].code, "22")
        self.assertEqual(rows[0]["base"], Decimal("18.00"))
        self.assertEqual(rows[0]["vat"], Decimal("3.96"))
        self.assertEqual(rows[1]["rate"].code, "10")


class PriceListTest(FlowTestBase):
    def test_adeguamento_percentuale_listino(self):
        pricelist = SupplierPriceList.objects.create(supplier=self.supplier, name="Listino 2026")
        item = PriceListItem.objects.create(pricelist=pricelist, product=self.product, price=Decimal("10.00"))

        apply_pricelist_adjustment(pricelist, Decimal("3"), user=self.user)
        item.refresh_from_db()
        self.assertEqual(item.price, Decimal("10.30"))
        self.assertEqual(pricelist.adjustments.count(), 1)
        self.assertEqual(pricelist.adjustments.first().items_count, 1)

    def test_prezzo_acquisto_da_listino(self):
        pricelist = SupplierPriceList.objects.create(supplier=self.supplier, name="Listino attivo")
        PriceListItem.objects.create(pricelist=pricelist, product=self.product, price=Decimal("4.00"), discount_pct=Decimal("10"))
        # Prezzo effettivo: 3.60; il prezzo articolo è 5.00 → vince il listino
        self.assertEqual(self.product.purchase_unit_price(self.supplier), Decimal("3.6000"))


class PagesSmokeTest(FlowTestBase):
    """Verifica che le pagine principali rispondano senza errori."""

    def setUp(self):
        self.login()

    def test_pagine_principali(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(product=self.product, description="x", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        quote.recalculate()
        order = SalesOrder.objects.create(customer=self.customer)
        order.lines.create(product=self.product, description="x", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        order.recalculate()
        po = PurchaseOrder.objects.create(supplier=self.supplier)
        po.lines.create(product=self.product, description="x", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("5"), vat_rate=self.vat22)
        po.recalculate()
        pricelist = SupplierPriceList.objects.create(supplier=self.supplier, name="Listino smoke")
        PriceListItem.objects.create(pricelist=pricelist, product=self.product, price=Decimal("5.00"))
        quote_template = QuoteTemplate.objects.create(name="Modello smoke")
        quote_template.lines.create(product=self.product, description="x", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        register_movement(product=self.product, delta=Decimal("5"), movement_type="load", user=self.user)

        urls = [
            reverse("core:home"),
            reverse("core:settings"),
            reverse("core:company_update"),
            reverse("core:vat_list"),
            reverse("core:vat_create"),
            reverse("core:uom_list"),
            reverse("core:tag_list"),
            reverse("core:paymentterm_list"),
            reverse("core:sequence_list"),
            reverse("accounts:user_list"),
            reverse("accounts:user_create"),
            reverse("accounts:password_change"),
            reverse("contacts:list"),
            reverse("contacts:create"),
            reverse("contacts:detail", args=[self.customer.pk]),
            reverse("contacts:update", args=[self.customer.pk]),
            reverse("catalog:product_list"),
            reverse("catalog:product_create"),
            reverse("catalog:product_import"),
            reverse("catalog:product_import_template"),
            reverse("catalog:product_detail", args=[self.product.pk]),
            reverse("catalog:product_update", args=[self.product.pk]),
            reverse("catalog:category_list"),
            reverse("inventory:stock_list"),
            reverse("inventory:movement_list"),
            reverse("inventory:adjust_create"),
            reverse("sales:quote_list"),
            reverse("sales:quote_create"),
            reverse("sales:quote_detail", args=[quote.pk]),
            reverse("sales:quote_update", args=[quote.pk]),
            reverse("sales:quote_print", args=[quote.pk]),
            reverse("sales:quote_template_list"),
            reverse("sales:quote_template_create"),
            reverse("sales:quote_template_update", args=[quote_template.pk]),
            reverse("sales:quote_template_data", args=[quote_template.pk]),
            reverse("sales:order_list"),
            reverse("sales:order_create"),
            reverse("sales:order_detail", args=[order.pk]),
            reverse("sales:order_update", args=[order.pk]),
            reverse("sales:order_print", args=[order.pk]),
            reverse("purchasing:po_list"),
            reverse("purchasing:po_create"),
            reverse("purchasing:po_detail", args=[po.pk]),
            reverse("purchasing:po_update", args=[po.pk]),
            reverse("purchasing:po_print", args=[po.pk]),
            reverse("purchasing:po_receive", args=[po.pk]),
            reverse("purchasing:pricelist_list"),
            reverse("purchasing:pricelist_create"),
            reverse("purchasing:pricelist_detail", args=[pricelist.pk]),
            reverse("purchasing:pricelist_update", args=[pricelist.pk]),
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200, f"{url} → {response.status_code}")

    def test_ricerca_drilldown(self):
        urls = [
            reverse("contacts:list") + "?q=Cliente&tipo=clienti",
            reverse("catalog:product_list") + "?sotto_scorta=1",
            reverse("inventory:stock_list") + "?sotto_scorta=1",
            reverse("sales:quote_list") + "?stato=aperti",
            reverse("sales:order_list") + "?stato=aperti",
            reverse("purchasing:po_list") + "?stato=aperti",
            reverse("purchasing:pricelist_list") + "?inattivi=1",
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)


class PermissionsTest(FlowTestBase):
    def test_anonimo_rediretto_al_login(self):
        urls = [
            reverse("core:home"),
            reverse("contacts:list"),
            reverse("catalog:product_list"),
            reverse("sales:quote_list"),
            reverse("inventory:stock_list"),
            reverse("catalog:product_defaults", args=[self.product.pk]),
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertIn(reverse("accounts:login"), response.url)

    def test_utente_senza_ruolo_non_accede(self):
        from django.contrib.auth.models import Group

        user = User.objects.create_user("mario", password="password123!")
        user.groups.add(Group.objects.get(name="Vendite"))
        self.client.force_login(user)

        # I preventivi sono accessibili al ruolo Vendite
        self.assertEqual(self.client.get(reverse("sales:quote_list")).status_code, 200)
        # Le impostazioni no
        self.assertEqual(self.client.get(reverse("core:settings")).status_code, 403)
        # Gli utenti no
        self.assertEqual(self.client.get(reverse("accounts:user_list")).status_code, 403)
        # I listini no
        self.assertEqual(self.client.get(reverse("purchasing:pricelist_list")).status_code, 403)


class QuoteEmailTest(FlowTestBase):
    """Invio del preventivo via email e visibilità di chi l'ha creato."""

    def setUp(self):
        self.login()
        self.quote = Quote.objects.create(customer=self.customer, created_by=self.user)
        self.quote.lines.create(
            product=self.product, description=self.product.name, qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("10.00"), vat_rate=self.vat22,
        )
        self.quote.recalculate()

    def test_pagina_mostra_il_modulo_email(self):
        response = self.client.get(reverse("sales:quote_detail", args=[self.quote.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("sales:quote_email", args=[self.quote.pk]))

    def test_invio_senza_email_configurata_avvisa(self):
        """Senza EMAIL_HOST l'invio non deve rompersi ma spiegare cosa manca."""
        response = self.client.post(
            reverse("sales:quote_email", args=[self.quote.pk]),
            {"to": "cliente@example.com", "subject": "Preventivo", "message": "Ciao"},
            follow=True,
        )
        self.assertContains(response, "Invio email non configurato")
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.status, Quote.STATUS_DRAFT, "lo stato non deve cambiare")

    def test_invio_senza_indirizzo_email(self):
        self.customer.email = ""
        self.customer.save(update_fields=["email"])
        response = self.client.post(
            reverse("sales:quote_email", args=[self.quote.pk]),
            {"to": "", "subject": "x", "message": "y"},
            follow=True,
        )
        self.assertContains(response, "indirizzo email del cliente")

    @override_settings(EMAIL_IS_CONFIGURED=True)
    def test_invio_riuscito_porta_il_preventivo_a_inviato(self):
        with mock.patch("apps.sales.emailing.send_quote_email", return_value=True) as invio:
            response = self.client.post(
                reverse("sales:quote_email", args=[self.quote.pk]),
                {"to": "cliente@example.com", "subject": "Preventivo X", "message": "Buongiorno"},
                follow=True,
            )
        self.assertEqual(invio.call_count, 1)
        self.assertContains(response, "inviato a cliente@example.com")
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.status, Quote.STATUS_SENT)

    @override_settings(EMAIL_IS_CONFIGURED=True)
    def test_invio_fallito_avvisa_e_non_cambia_stato(self):
        with mock.patch("apps.sales.emailing.send_quote_email", side_effect=Exception("SMTP non raggiungibile")):
            response = self.client.post(
                reverse("sales:quote_email", args=[self.quote.pk]),
                {"to": "cliente@example.com", "subject": "x", "message": "y"},
                follow=True,
            )
        self.assertContains(response, "Invio non riuscito")
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.status, Quote.STATUS_DRAFT)

    def test_invio_su_preventivo_gia_inviato_non_cambia_stato(self):
        self.quote.status = Quote.STATUS_ACCEPTED
        self.quote.save(update_fields=["status"])
        with mock.patch("apps.sales.emailing.send_quote_email", return_value=True):
            self.client.post(
                reverse("sales:quote_email", args=[self.quote.pk]),
                {"to": "cliente@example.com", "subject": "x", "message": "y"},
            )
        self.quote.refresh_from_db()
        self.assertEqual(self.quote.status, Quote.STATUS_ACCEPTED)


class QuoteCreatorTest(FlowTestBase):
    """L'elenco deve mostrare chi ha creato il preventivo e permettere il filtro."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.altro = User.objects.create_user("lucia", password="password123!", first_name="Lucia", last_name="Bianchi")
        cls.mio = Quote.objects.create(customer=cls.customer, created_by=cls.user, reference="Mio")
        cls.suo = Quote.objects.create(customer=cls.customer, created_by=cls.altro, reference="Suo")

    def setUp(self):
        self.login()

    def test_elenco_mostra_chi_ha_creato(self):
        response = self.client.get(reverse("sales:quote_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "admin")
        self.assertContains(response, "Lucia Bianchi")

    def test_filtro_per_utente(self):
        response = self.client.get(reverse("sales:quote_list"), {"utente": self.altro.pk})
        quotes = list(response.context["quotes"])
        self.assertEqual(len(quotes), 1)
        self.assertEqual(quotes[0].pk, self.suo.pk)

    def test_filtro_solo_i_miei(self):
        response = self.client.get(reverse("sales:quote_list"), {"miei": "1"})
        quotes = list(response.context["quotes"])
        self.assertEqual(len(quotes), 1)
        self.assertEqual(quotes[0].pk, self.mio.pk)

    def test_elenco_creatori_contiene_solo_chi_ha_preventivi(self):
        response = self.client.get(reverse("sales:quote_list"))
        nomi = {u.username for u in response.context["creators"]}
        self.assertIn("admin", nomi)
        self.assertIn("lucia", nomi)


class QuotePdfTest(FlowTestBase):
    """Il PDF del preventivo deve essere generato quando WeasyPrint è disponibile."""

    def setUp(self):
        self.login()
        self.quote = Quote.objects.create(customer=self.customer)
        self.quote.lines.create(
            product=self.product, description=self.product.name, qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("10.00"), vat_rate=self.vat22,
        )
        self.quote.recalculate()

    def test_pdf_generato_oppure_assente_ma_senza_errori(self):
        from apps.core.pdf import pdf_available
        from apps.sales.pdf import render_quote_pdf

        pdf = render_quote_pdf(self.quote)
        if pdf_available():
            self.assertIsNotNone(pdf, "con WeasyPrint disponibile il PDF deve essere generato")
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertGreater(len(pdf), 1000)
        else:
            self.assertIsNone(pdf, "senza WeasyPrint il PDF e' None e il gestionale continua a funzionare")


class CommissionTest(FlowTestBase):
    """La provvigione a chi ha presentato il cliente erode il margine."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.beneficiario = Contact.objects.create(
            name="Geom. Bianchi", is_customer=False, is_supplier=False, city="Milano"
        )
        # secondo articolo, per verificare la ripartizione della provvigione
        cls.altro_prodotto = Product.objects.create(
            name="Vaso sospeso",
            uom=cls.uom,
            sale_price=Decimal("30.00"),
            sale_vat=cls.vat22,
            purchase_price=Decimal("6.00"),
            purchase_vat=cls.vat22,
            main_supplier=cls.supplier,
        )

    def _ordine_consegnato(self, commission_pct="10", commission_contact=None):
        """Ordine con due righe: 100 di imponibile a costo 40, 300 a costo 60."""
        order = SalesOrder.objects.create(
            customer=self.customer,
            status=SalesOrder.STATUS_DELIVERED,
            delivered_at=timezone.now(),
            commission_contact=self.beneficiario if commission_pct else None if commission_contact is None else commission_contact,
            commission_pct=Decimal(commission_pct or "0"),
        )
        order.lines.create(
            product=self.product, description=self.product.name, qty=Decimal("10"),
            uom=self.uom, unit_price=Decimal("10.00"), vat_rate=self.vat22,
            qty_delivered=Decimal("10"), unit_cost=Decimal("4.00"),
        )
        order.lines.create(
            product=self.altro_prodotto, description=self.altro_prodotto.name, qty=Decimal("10"),
            uom=self.uom, unit_price=Decimal("30.00"), vat_rate=self.vat22,
            qty_delivered=Decimal("10"), unit_cost=Decimal("6.00"),
        )
        order.recalculate()
        return order

    def test_importo_provvigione(self):
        order = self._ordine_consegnato(commission_pct="10")
        self.assertTrue(order.has_commission)
        # 10% di 400 di imponibile
        self.assertEqual(order.commission_amount, Decimal("40.00"))

    def test_provvigione_zero_senza_beneficiario(self):
        order = SalesOrder.objects.create(customer=self.customer, commission_pct=Decimal("10"))
        self.assertFalse(order.has_commission)
        self.assertEqual(order.commission_amount, Decimal("0"))

    def test_provvigione_zero_senza_percentuale(self):
        order = SalesOrder.objects.create(customer=self.customer, commission_contact=self.beneficiario)
        self.assertFalse(order.has_commission)

    def test_provvigione_erode_il_margine(self):
        self._ordine_consegnato(commission_pct="10")
        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["revenue"], Decimal("400.00"))
        self.assertEqual(summary["cost"], Decimal("100.00"))
        self.assertEqual(summary["commission"], Decimal("40.00"))
        # 400 - 100 - 40
        self.assertEqual(summary["margin"], Decimal("260.00"))
        self.assertEqual(summary["margin_pct"], Decimal("65.00"))

    def test_senza_provvigione_il_margine_non_cambia(self):
        self._ordine_consegnato(commission_pct="0")
        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["commission"], Decimal("0"))
        self.assertEqual(summary["margin"], Decimal("300.00"))

    def test_margine_prima_della_provvigione_resta_disponibile(self):
        self._ordine_consegnato(commission_pct="10")
        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["margin_before_commission"], Decimal("300.00"))
        self.assertEqual(
            summary["margin_before_commission"] - summary["commission"], summary["margin"]
        )

    def test_provvigione_ripartita_sulle_righe_in_proporzione(self):
        self._ordine_consegnato(commission_pct="10")
        per_articolo = {row["label"]: row for row in analytics.by_product(analytics.PERIOD_YEAR)}
        # riga da 100 -> 10 di provvigione, riga da 300 -> 30
        self.assertEqual(per_articolo["Bullone M8"]["commission"], Decimal("10.00"))
        self.assertEqual(per_articolo["Vaso sospeso"]["commission"], Decimal("30.00"))
        self.assertEqual(per_articolo["Bullone M8"]["margin"], Decimal("50.00"))  # 100 - 40 - 10
        self.assertEqual(per_articolo["Vaso sospeso"]["margin"], Decimal("210.00"))  # 300-60-30

        # e la somma delle quote deve tornare col totale
        totale = sum(row["commission"] for row in analytics.by_product(analytics.PERIOD_YEAR))
        self.assertEqual(totale, Decimal("40.00"))

    def test_breakdowns_della_dashboard_include_la_provvigione(self):
        self._ordine_consegnato(commission_pct="10")
        stats = analytics.breakdowns(analytics.PERIOD_YEAR, limit=10)
        self.assertEqual(stats["summary"]["commission"], Decimal("40.00"))
        self.assertEqual(stats["summary"]["margin"], Decimal("260.00"))
        somma = sum(row["commission"] for row in stats["by_product"])
        self.assertEqual(somma, Decimal("40.00"))

    def test_provvigione_erode_anche_il_grafico_mensile(self):
        self._ordine_consegnato(commission_pct="10")
        serie = analytics.monthly_series(months=12)
        self.assertEqual(sum(serie["margin"]), 260.0)

    def test_la_provvigione_segue_il_preventivo_nell_ordine(self):
        self.login()
        quote = Quote.objects.create(
            customer=self.customer,
            commission_contact=self.beneficiario,
            commission_pct=Decimal("7.5"),
        )
        quote.lines.create(
            product=self.product, description=self.product.name, qty=Decimal("10"),
            uom=self.uom, unit_price=Decimal("10.00"), vat_rate=self.vat22,
        )
        quote.recalculate()
        quote.status = Quote.STATUS_ACCEPTED
        quote.save(update_fields=["status"])

        response = self.client.post(reverse("sales:quote_convert", args=[quote.pk]))
        self.assertEqual(response.status_code, 302)
        order = SalesOrder.objects.get(source_quote=quote)
        self.assertEqual(order.commission_contact, self.beneficiario)
        self.assertEqual(order.commission_pct, Decimal("7.5"))
        self.assertEqual(order.commission_amount, Decimal("7.50"))  # 7,5% di 100

    def test_il_form_del_preventivo_mostra_i_campi_provvigione(self):
        self.login()
        response = self.client.get(reverse("sales:quote_create"))
        self.assertContains(response, 'name="commission_contact"')
        self.assertContains(response, 'name="commission_pct"')

    def test_le_fatture_non_hanno_la_provvigione(self):
        """La provvigione non deve comparire su documenti di acquisto o fatture."""
        from apps.billing.models import PurchaseInvoice, SalesInvoice
        from apps.purchasing.models import PurchaseOrder

        for modello in (SalesInvoice, PurchaseInvoice, PurchaseOrder):
            with self.subTest(modello=modello.__name__):
                self.assertFalse(hasattr(modello, "commission_pct"))


class ExportExcelTest(FlowTestBase):
    """Esportazione Excel del dettaglio vendite."""

    def setUp(self):
        self.login()

    def _ordine(self, qty="10", price="10.00", consegnato=True, con_provvigione=False):
        order = SalesOrder.objects.create(
            customer=self.customer,
            status=SalesOrder.STATUS_DELIVERED if consegnato else SalesOrder.STATUS_CONFIRMED,
            delivered_at=timezone.now() if consegnato else None,
            created_by=self.user,
            commission_contact=self.customer if con_provvigione else None,
            commission_pct=Decimal("10") if con_provvigione else Decimal("0"),
        )
        order.lines.create(
            product=self.product, description=self.product.name, qty=Decimal(qty),
            uom=self.uom, unit_price=Decimal(price), vat_rate=self.vat22,
            qty_delivered=Decimal(qty), unit_cost=Decimal("4.00"),
        )
        order.recalculate()
        return order

    def _foglio(self, **parametri):
        import openpyxl
        from io import BytesIO

        response = self.client.get(reverse("sales:export_sales_excel"), parametri)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        return openpyxl.load_workbook(BytesIO(response.content)).active

    def test_file_valido_con_intestazioni_chiare(self):
        self._ordine()
        foglio = self._foglio(periodo="anno")
        intestazioni = [cella.value for cella in foglio[1]]
        self.assertEqual(intestazioni[0], "Data")
        self.assertEqual(intestazioni[2], "Cliente")
        self.assertEqual(intestazioni[3], "Venditore")
        self.assertIn("Provvigione", intestazioni)
        self.assertIn("Margine", intestazioni)
        self.assertEqual(foglio.max_column, 15)

    def test_una_riga_per_articolo_venduto(self):
        self._ordine()
        foglio = self._foglio(periodo="anno")
        # 1 riga di intestazione + 1 riga di dati + 1 riga totali
        self.assertEqual(foglio.max_row, 3)

    def test_esclude_gli_ordini_non_consegnati(self):
        self._ordine(consegnato=False)
        foglio = self._foglio(periodo="anno")
        self.assertEqual(foglio.max_row, 1, "solo l'intestazione: nessuna vendita consegnata")

    def test_riporta_i_valori_e_i_totali(self):
        ordine = self._ordine(qty="10", price="10.00")
        foglio = self._foglio(periodo="anno")
        riga = [cella.value for cella in foglio[2]]
        self.assertEqual(riga[1], ordine.number, "deve riportare il numero del documento")
        self.assertEqual(riga[2], self.customer.name, "deve riportare il cliente")
        self.assertEqual(riga[5], "Bullone M8")
        self.assertEqual(riga[6], 10.0)          # quantità
        self.assertEqual(riga[7], "PZ")          # unità di misura
        self.assertEqual(riga[8], 10.0)          # prezzo
        self.assertEqual(riga[10], 100.0)        # imponibile
        self.assertEqual(riga[11], 40.0)         # costo
        self.assertEqual(riga[13], 60.0)         # margine

        totali = [cella.value for cella in foglio[foglio.max_row]]
        self.assertEqual(totali[5], "TOTALE")
        self.assertEqual(totali[10], 100.0)
        self.assertEqual(totali[11], 40.0)
        self.assertEqual(totali[13], 60.0)

    def test_la_provvigione_compare_e_riduce_il_margine(self):
        self._ordine(con_provvigione=True)
        foglio = self._foglio(periodo="anno")
        riga = [cella.value for cella in foglio[2]]
        self.assertEqual(riga[12], 10.0, "provvigione 10% di 100")
        self.assertEqual(riga[13], 50.0, "margine 100 - 40 - 10")
        self.assertEqual(riga[14], 50.0, "margine %")

    def test_il_venditore_e_riportato(self):
        self._ordine()
        foglio = self._foglio(periodo="anno")
        riga = [cella.value for cella in foglio[2]]
        self.assertEqual(riga[3], "admin")

    def test_la_data_e_un_vero_formato_data(self):
        import datetime

        self._ordine()
        foglio = self._foglio(periodo="anno")
        valore = foglio.cell(row=2, column=1).value
        self.assertIsInstance(valore, (datetime.date, datetime.datetime))
        self.assertEqual(foglio.cell(row=2, column=1).number_format, "DD/MM/YYYY")

    def test_periodo_non_valido_ricade_sull_anno(self):
        self._ordine()
        foglio = self._foglio(periodo="qualcosa-di-strano")
        self.assertEqual(foglio.max_row, 3)


class DiscountAndRoundingTest(FlowTestBase):
    """Lo sconto di riga e il modo di arrotondare gli importi."""

    def _riga(self, qty, price, discount, vat=None):
        quote = Quote.objects.create(customer=self.customer)
        riga = QuoteLine.objects.create(
            quote=quote, position=1, description="prova", qty=Decimal(qty), uom=self.uom,
            unit_price=Decimal(price), discount_pct=Decimal(discount), vat_rate=vat or self.vat22,
        )
        return quote, riga

    def test_lo_sconto_riduce_l_imponibile_della_riga(self):
        for qty, price, discount, atteso in [
            ("10", "100.00", "0", "1000.00"),
            ("10", "100.00", "10", "900.00"),
            ("10", "100.00", "50", "500.00"),
            ("10", "100.00", "100", "0.00"),
            ("1", "100.00", "33.33", "66.67"),
            ("3", "33.3333", "10", "90.00"),
            ("0.5", "1000.00", "15", "425.00"),
        ]:
            with self.subTest(qty=qty, sconto=discount):
                _, riga = self._riga(qty, price, discount)
                self.assertEqual(riga.line_subtotal, Decimal(atteso))

    def test_lo_sconto_si_riflette_sui_totali_del_documento(self):
        quote, _ = self._riga("10", "100.00", "20")
        quote.recalculate()
        quote.refresh_from_db()
        self.assertEqual(quote.subtotal, Decimal("800.00"))
        self.assertEqual(quote.vat_total, Decimal("176.00"))  # 22% di 800
        self.assertEqual(quote.grand_total, Decimal("976.00"))

    def test_l_iva_si_calcola_sull_imponibile_scontato(self):
        _, riga = self._riga("10", "100.00", "50")
        self.assertEqual(riga.line_subtotal, Decimal("500.00"))
        self.assertEqual(riga.line_vat, Decimal("110.00"))  # 22% di 500, non di 1000

    def test_lo_sconto_arriva_al_fatturato(self):
        quote, _ = self._riga("10", "100.00", "20")
        quote.status = Quote.STATUS_ACCEPTED
        quote.save(update_fields=["status"])
        order = services.convert_quote_to_order(quote)
        order.status = SalesOrder.STATUS_DELIVERED
        order.delivered_at = timezone.now()
        order.save(update_fields=["status", "delivered_at"])

        summary = analytics.summary(analytics.PERIOD_YEAR)
        self.assertEqual(summary["revenue"], Decimal("800.00"), "il fatturato deve tenere conto dello sconto")

    def test_i_totali_tornano_con_piu_righe_scontate(self):
        quote = Quote.objects.create(customer=self.customer)
        QuoteLine.objects.create(
            quote=quote, position=1, description="a", qty=Decimal("3"), uom=self.uom,
            unit_price=Decimal("10.00"), discount_pct=Decimal("10"), vat_rate=self.vat22,
        )
        QuoteLine.objects.create(
            quote=quote, position=2, description="b", qty=Decimal("2"), uom=self.uom,
            unit_price=Decimal("20.00"), discount_pct=Decimal("5"), vat_rate=self.vat10,
        )
        quote.recalculate()
        quote.refresh_from_db()
        # 27,00 + 38,00
        self.assertEqual(quote.subtotal, Decimal("65.00"))
        # 22% di 27 = 5,94  +  10% di 38 = 3,80
        self.assertEqual(quote.vat_total, Decimal("9.74"))
        self.assertEqual(quote.grand_total, Decimal("74.74"))

    def test_arrotondamento_commerciale_non_bancario(self):
        """Con l'arrotondamento bancario 0,665 diventerebbe 0,66."""
        from apps.purchasing.models import round4
        from apps.sales.models import round2, round3

        for valore, atteso in [("0.665", "0.67"), ("1.005", "1.01"), ("0.125", "0.13"), ("2.675", "2.68")]:
            with self.subTest(valore=valore, funzione="round2"):
                self.assertEqual(round2(Decimal(valore)), Decimal(atteso))
        self.assertEqual(round3(Decimal("1.0005")), Decimal("1.001"))
        self.assertEqual(round4(Decimal("1.00005")), Decimal("1.0001"))

    def test_i_totali_mostrati_in_pagina_tornano_col_calcolo(self):
        self.login()
        response = self.client.post(
            reverse("sales:quote_create"),
            {
                "customer": self.customer.pk,
                "date": "2026-01-10",
                "valid_until": "",
                "payment_term": "",
                "reference": "",
                "commission_contact": "",
                "commission_pct": "0",
                "terms_text": "",
                "notes": "",
                "lines-TOTAL_FORMS": "1",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-product": self.product.pk,
                "lines-0-description": "",
                "lines-0-qty": "10",
                "lines-0-uom": self.uom.pk,
                "lines-0-unit_price": "100.00",
                "lines-0-discount_pct": "25",
                "lines-0-vat_rate": self.vat22.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        quote = Quote.objects.latest("pk")
        self.assertEqual(quote.subtotal, Decimal("750.00"), "lo sconto inserito nel form deve essere salvato")
        self.assertEqual(quote.lines.get().discount_pct, Decimal("25.00"))


class BugVariTest(FlowTestBase):
    """Regressioni su difetti visti nella creazione dei preventivi."""

    def setUp(self):
        self.login()

    # ------------------------------------------------ commenti visibili in pagina
    def test_nessun_commento_su_piu_righe_nei_template(self):
        """{# #} in Django vale solo su una riga: su più righe finisce in pagina."""
        import re
        from pathlib import Path

        radice = Path(__file__).resolve().parent.parent.parent / "templates"
        trovati = []
        for percorso in radice.rglob("*.html"):
            testo = percorso.read_text(encoding="utf-8")
            for commento in re.finditer(r"\{#(.*?)#\}", testo, re.DOTALL):
                if "\n" in commento.group(1):
                    trovati.append(f"{percorso.name}:{testo[: commento.start()].count(chr(10)) + 1}")
        self.assertEqual(trovati, [], "commenti {# #} su più righe: verrebbero mostrati all'utente")

    def test_la_pagina_non_mostra_commenti(self):
        response = self.client.get(reverse("sales:quote_create"))
        contenuto = response.content.decode("utf-8")
        self.assertNotIn("{#", contenuto)
        self.assertNotIn("#}", contenuto)

    # ------------------------------------------- righe stile Odoo (sezioni)
    def test_intestazioni_e_pulsanti_delle_righe(self):
        """Le sezioni sono righe dedicate: niente più colonna «Sezione»."""
        response = self.client.get(reverse("sales:quote_create"))
        pagina = response.content.decode("utf-8")

        # ordine dichiarato nelle intestazioni
        intestazioni = [m.group(1) for m in re.finditer(r"<th[^>]*>(.*?)</th>", pagina, re.DOTALL)]
        testi = [re.sub(r"<[^>]+>", "", t).strip() for t in intestazioni]
        self.assertIn("Descrizione", testi)
        self.assertNotIn("Sezione", testi, "la colonna «Sezione» è sostituita dalle righe dedicate")

        # pulsanti per aggiungere riga, sezione, sottosezione e nota
        for tipo in ("article", "section", "subsection", "note"):
            self.assertIn(f'data-add-line-type="{tipo}"', pagina)

        # ordine effettivo dei campi nella prima riga (nomi completi, per non
        # confondersi con «nav-section» della barra laterale)
        posizione_descrizione = pagina.index('name="lines-0-description"')
        posizione_qta = pagina.index('name="lines-0-qty"')
        self.assertLess(
            posizione_descrizione, posizione_qta,
            "il campo descrizione deve venire prima di quello della quantità",
        )
        # posizione e tipo viaggiano nascosti: l'ordine lo decide l'utente
        self.assertIn('name="lines-0-position"', pagina)
        self.assertIn('name="lines-0-line_type"', pagina)
        self.assertIn(
            'class="btn btn-outline-secondary btn-sm move-line"', pagina,
            "le frecce su/giù devono comparire nelle righe del preventivo",
        )

    # ------------------------------------------------ sconto non in stampa
    def _preventivo_scontato(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(
            product=self.product, description="Con sconto", qty=Decimal("10"),
            uom=self.uom, unit_price=Decimal("100.00"), discount_pct=Decimal("20"),
            vat_rate=self.vat22,
        )
        quote.recalculate()
        return quote

    def test_il_preventivo_stampato_non_mostra_lo_sconto(self):
        quote = self._preventivo_scontato()
        response = self.client.get(reverse("sales:quote_print", args=[quote.pk]))
        pagina = response.content.decode("utf-8")
        self.assertNotIn("Sc. %", pagina, "lo sconto non deve comparire nel preventivo per il cliente")
        self.assertNotIn("20%", pagina, "la percentuale di sconto non deve comparire")

    def test_il_preventivo_stampato_mostra_il_prezzo_gia_scontato(self):
        quote = self._preventivo_scontato()
        response = self.client.get(reverse("sales:quote_print", args=[quote.pk]))
        pagina = response.content.decode("utf-8")
        # 100 meno il 20% = 80,00: è questo che deve leggere il cliente
        self.assertIn("80,00", pagina, "deve comparire il prezzo netto, non quello di listino")
        self.assertNotIn("100,00", pagina, "il prezzo di listino non deve comparire")

    def test_il_totale_del_preventivo_stampato_e_quello_scontato(self):
        quote = self._preventivo_scontato()
        response = self.client.get(reverse("sales:quote_print", args=[quote.pk]))
        pagina = response.content.decode("utf-8")
        # imponibile 800, IVA 176, totale 976
        self.assertIn("800,00", pagina)
        self.assertIn("976,00", pagina)

    def test_la_fattura_continua_a_mostrare_lo_sconto(self):
        """Alla fattura non è stato chiesto: resta com'era."""
        from apps.billing.models import SalesInvoice

        fattura = SalesInvoice.objects.create(customer=self.customer, date=timezone.localdate())
        fattura.lines.create(
            product=self.product, description="Riga", qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("50.00"), discount_pct=Decimal("10"),
            vat_rate=self.vat22,
        )
        fattura.recalculate()
        response = self.client.get(reverse("billing:salesinvoice_print", args=[fattura.pk]))
        pagina = response.content.decode("utf-8")
        self.assertIn("Sc. %", pagina, "la fattura mostra ancora lo sconto")

    def test_le_colonne_della_stampa_sono_coerenti(self):
        """Con lo sconto nascosto i colspan delle sezioni devono tornare."""
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(section="Bagno 1", product=self.product, description="a", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22)
        quote.lines.create(section="Bagno 1", product=self.product, description="b", qty=Decimal("1"), uom=self.uom, unit_price=Decimal("5"), vat_rate=self.vat22)
        quote.recalculate()
        response = self.client.get(reverse("sales:quote_print", args=[quote.pk]))
        pagina = response.content.decode("utf-8")
        self.assertIn("Totale Bagno 1", pagina)
        # senza la colonna sconto le colonne sono 7: il totale di sezione ne occupa 6
        self.assertIn('colspan="7"', pagina, "la riga di sezione deve occupare 7 colonne")
        self.assertIn('colspan="6"', pagina, "il totale di sezione deve occuparne 6 più il valore")

    def test_il_prezzo_netto_arrotonda_bene_i_casi_bassi(self):
        """0,35 scontato del 5% fa 0,3325: al cliente serve il valore esatto."""
        from apps.core.templatetags.core_extras import net_price

        self.assertEqual(net_price(Decimal("0.3325")), "0,3325")
        self.assertEqual(net_price(Decimal("100")), "100,00")
        self.assertEqual(net_price(Decimal("1234.5")), "1.234,50")

    def test_prezzo_netto_della_riga(self):
        quote = Quote.objects.create(customer=self.customer)
        riga = quote.lines.create(
            product=self.product, description="x", qty=Decimal("1"), uom=self.uom,
            unit_price=Decimal("100.00"), discount_pct=Decimal("20"), vat_rate=self.vat22,
        )
        self.assertEqual(riga.net_unit_price, Decimal("80.0000"))
        riga.discount_pct = Decimal("0")
        self.assertEqual(riga.net_unit_price, Decimal("100.0000"))


class RigheStileOdooTest(FlowTestBase):
    """Sezioni, sottosezioni e note (righe di testo) e spostamento delle righe."""

    def setUp(self):
        self.login()

    def test_sezioni_e_note_non_entrano_nei_totali(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(line_type="section", description="Bagno padronale")
        quote.lines.create(line_type="subsection", description="Mobile lavabo")
        quote.lines.create(line_type="note", description="modello: MODO PROJECT")
        quote.lines.create(
            product=self.product, description="a", qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        quote.lines.create(line_type="section", description="Cucina")
        quote.lines.create(
            product=self.product, description="b", qty=Decimal("1"),
            uom=self.uom, unit_price=Decimal("100"), vat_rate=self.vat22,
        )
        quote.recalculate()
        quote.refresh_from_db()

        self.assertEqual(quote.subtotal, Decimal("120.00"), "le righe di testo non devono pesare nei totali")

        gruppi = group_lines_by_section(quote.lines.order_by("position", "pk"))
        self.assertEqual([g["section"] for g in gruppi], ["Bagno padronale", "Cucina"])
        self.assertEqual(gruppi[0]["subtotal"], Decimal("20.00"))
        self.assertEqual(gruppi[1]["subtotal"], Decimal("100.00"))

        # la sottosezione apre un blocco con il subtotale delle sole righe sotto di lei
        blocchi = gruppi[0]["blocks"]
        self.assertEqual([b["subsection"] for b in blocchi], ["", "Mobile lavabo"])
        self.assertEqual(blocchi[1]["subtotal"], Decimal("20.00"))

        note = [l for l in blocchi[1]["lines"] if l.line_type == "note"]
        self.assertEqual(len(note), 1)
        self.assertIsNone(note[0].row_number)
        articoli = [l for l in blocchi[1]["lines"] if l.line_type == "article"]
        self.assertEqual(articoli[0].row_number, 1)

    def test_subtotale_per_sottosezione(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(line_type="section", description="Bagno")
        quote.lines.create(line_type="subsection", description="Mobile")
        quote.lines.create(
            product=self.product, description="a", qty=Decimal("1"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        quote.lines.create(line_type="subsection", description="Sanitari")
        quote.lines.create(
            product=self.product, description="b", qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        gruppi = group_lines_by_section(quote.lines.order_by("position", "pk"))
        self.assertEqual([g["section"] for g in gruppi], ["Bagno"])
        blocchi = gruppi[0]["blocks"]
        self.assertEqual([b["subsection"] for b in blocchi], ["", "Mobile", "Sanitari"])
        self.assertEqual(blocchi[1]["subtotal"], Decimal("10.00"))
        self.assertEqual(blocchi[2]["subtotal"], Decimal("20.00"))
        self.assertEqual(gruppi[0]["subtotal"], Decimal("30.00"))

    def test_la_stampa_mostra_i_subtotali_delle_sottosezioni(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(line_type="section", description="Bagno")
        quote.lines.create(line_type="subsection", description="Mobile")
        quote.lines.create(
            product=self.product, description="a", qty=Decimal("1"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        quote.lines.create(line_type="subsection", description="Sanitari")
        quote.lines.create(
            product=self.product, description="b", qty=Decimal("2"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        quote.recalculate()
        response = self.client.get(reverse("sales:quote_print", args=[quote.pk]))
        pagina = response.content.decode("utf-8")
        self.assertIn("Totale Mobile", pagina)
        self.assertIn("Totale Sanitari", pagina)
        self.assertIn("Totale Bagno", pagina)

    def test_la_riga_vuota_di_servizio_non_blocca_il_salvataggio(self):
        """Il modulo porta sempre una riga vuota in più: il JavaScript le assegna
        una posizione, ma non deve diventare una riga da validare e salvare."""
        response = self.client.post(
            reverse("sales:quote_create"),
            {
                "customer": self.customer.pk,
                "date": "2026-01-10",
                "valid_until": "",
                "payment_term": "",
                "reference": "",
                "commission_contact": "",
                "commission_pct": "0",
                "terms_text": "",
                "notes": "",
                "lines-TOTAL_FORMS": "6",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-line_type": "section",
                "lines-0-position": "1",
                "lines-0-description": "Bagno padronale",
                "lines-0-qty": "1",
                "lines-0-unit_price": "0",
                "lines-0-discount_pct": "0",
                "lines-1-line_type": "subsection",
                "lines-1-position": "2",
                "lines-1-description": "Mobile lavabo",
                "lines-1-qty": "1",
                "lines-1-unit_price": "0",
                "lines-1-discount_pct": "0",
                "lines-2-line_type": "article",
                "lines-2-position": "3",
                "lines-2-product": self.product.pk,
                "lines-2-description": "Mobile",
                "lines-2-qty": "1",
                "lines-2-uom": self.uom.pk,
                "lines-2-unit_price": "100",
                "lines-2-discount_pct": "0",
                "lines-2-vat_rate": self.vat22.pk,
                "lines-3-line_type": "subsection",
                "lines-3-position": "4",
                "lines-3-description": "Sanitari",
                "lines-3-qty": "1",
                "lines-3-unit_price": "0",
                "lines-3-discount_pct": "0",
                "lines-4-line_type": "article",
                "lines-4-position": "5",
                "lines-4-product": self.product.pk,
                "lines-4-description": "Sanitari",
                "lines-4-qty": "1",
                "lines-4-uom": self.uom.pk,
                "lines-4-unit_price": "50",
                "lines-4-discount_pct": "0",
                "lines-4-vat_rate": self.vat22.pk,
                # riga vuota di servizio, rinumerata dal JavaScript
                "lines-5-line_type": "article",
                "lines-5-position": "6",
                "lines-5-description": "",
                "lines-5-qty": "1",
                "lines-5-unit_price": "0",
                "lines-5-discount_pct": "0",
            },
        )
        self.assertEqual(response.status_code, 302, "il preventivo con più sottosezioni deve salvarsi")
        quote = Quote.objects.latest("pk")
        self.assertEqual(quote.lines.count(), 5, "la riga vuota di servizio non va salvata")
        self.assertEqual(quote.subtotal, Decimal("150.00"))

    def test_crea_preventivo_con_sezione_dal_modulo(self):
        response = self.client.post(
            reverse("sales:quote_create"),
            {
                "customer": self.customer.pk,
                "date": "2026-01-10",
                "valid_until": "",
                "payment_term": "",
                "reference": "",
                "commission_contact": "",
                "commission_pct": "0",
                "terms_text": "",
                "notes": "",
                "lines-TOTAL_FORMS": "2",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-line_type": "section",
                "lines-0-position": "1",
                "lines-0-description": "Bagno padronale",
                "lines-0-qty": "1",
                "lines-0-unit_price": "0",
                "lines-0-discount_pct": "0",
                "lines-1-line_type": "article",
                "lines-1-position": "2",
                "lines-1-product": self.product.pk,
                "lines-1-description": "Mobile lavabo",
                "lines-1-qty": "1",
                "lines-1-uom": self.uom.pk,
                "lines-1-unit_price": "100",
                "lines-1-discount_pct": "0",
                "lines-1-vat_rate": self.vat22.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        quote = Quote.objects.latest("pk")
        self.assertEqual(quote.lines.count(), 2)
        self.assertEqual(quote.subtotal, Decimal("100.00"))
        sezione = quote.lines.get(line_type="section")
        self.assertEqual(sezione.description, "Bagno padronale")
        self.assertEqual(sezione.position, 1)
        self.assertEqual(quote.lines.get(line_type="article").position, 2)

    def test_spostare_le_righe_cambia_l_ordine(self):
        quote = Quote.objects.create(customer=self.customer)
        prima = quote.lines.create(
            product=self.product, description="prima", qty=Decimal("1"), uom=self.uom,
            unit_price=Decimal("10"), vat_rate=self.vat22, position=1,
        )
        seconda = quote.lines.create(
            product=self.product, description="seconda", qty=Decimal("1"), uom=self.uom,
            unit_price=Decimal("5"), vat_rate=self.vat22, position=2,
        )
        quote.recalculate()

        # le frecce su/giù rimettono la posizione nei campi nascosti: al salvataggio
        # l'ordine visibile è quello che resta
        response = self.client.post(
            reverse("sales:quote_update", args=[quote.pk]),
            {
                "customer": self.customer.pk,
                "date": "2026-01-10",
                "valid_until": "",
                "payment_term": "",
                "reference": "",
                "commission_contact": "",
                "commission_pct": "0",
                "terms_text": "",
                "notes": "",
                "lines-TOTAL_FORMS": "2",
                "lines-INITIAL_FORMS": "2",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-id": prima.pk,
                "lines-0-position": "2",
                "lines-0-line_type": "article",
                "lines-0-description": "prima",
                "lines-0-qty": "1",
                "lines-0-uom": self.uom.pk,
                "lines-0-unit_price": "10",
                "lines-0-discount_pct": "0",
                "lines-0-vat_rate": self.vat22.pk,
                "lines-1-id": seconda.pk,
                "lines-1-position": "1",
                "lines-1-line_type": "article",
                "lines-1-description": "seconda",
                "lines-1-qty": "1",
                "lines-1-uom": self.uom.pk,
                "lines-1-unit_price": "5",
                "lines-1-discount_pct": "0",
                "lines-1-vat_rate": self.vat22.pk,
            },
        )
        self.assertEqual(response.status_code, 302)
        prima.refresh_from_db()
        seconda.refresh_from_db()
        self.assertEqual((prima.position, seconda.position), (2, 1))
        ordine = list(quote.lines.order_by("position", "pk").values_list("description", flat=True))
        self.assertEqual(ordine, ["seconda", "prima"])

    def test_convertire_in_ordine_mantiene_le_sezioni(self):
        quote = Quote.objects.create(customer=self.customer)
        quote.lines.create(line_type="section", description="Bagno padronale")
        quote.lines.create(
            product=self.product, description="a", qty=Decimal("1"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        quote.recalculate()
        order = services.convert_quote_to_order(quote, user=self.user)
        self.assertEqual(order.lines.get(line_type="section").description, "Bagno padronale")
        self.assertEqual(order.subtotal, Decimal("10.00"))

    def test_niente_frecce_nei_documenti_che_non_le_usano(self):
        """Gli ordini fornitore non inviano la posizione: niente frecce (e niente promesse)."""
        response = self.client.get(reverse("purchasing:po_create"))
        pagina = response.content.decode("utf-8")
        self.assertIn("lines-0-product", pagina)
        self.assertNotIn('class="btn btn-outline-secondary btn-sm move-line"', pagina)

    def test_il_modello_riporta_il_tipo_riga(self):
        template = QuoteTemplate.objects.create(name="Bagno tipo")
        template.lines.create(line_type="section", description="Bagno")
        template.lines.create(
            product=self.product, description="a", qty=Decimal("1"),
            uom=self.uom, unit_price=Decimal("10"), vat_rate=self.vat22,
        )
        response = self.client.get(reverse("sales:quote_template_data", args=[template.pk]))
        dati = response.json()
        self.assertEqual(dati["lines"][0]["line_type"], "section")
        self.assertEqual(dati["lines"][1]["line_type"], "article")
